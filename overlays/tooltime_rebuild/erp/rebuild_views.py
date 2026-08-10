from __future__ import annotations

import base64
import csv
import io
import json
import os
import re
from datetime import timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django import forms
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.files.base import ContentFile
from django.db import transaction
from django.db.models import Q, Sum
from django.http import JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.views.decorators.http import require_http_methods, require_POST

from . import models as m


# KAYI Next intentionally keeps the existing ERP data model. This module replaces
# the legacy interaction model with a ToolTime-like operational flow while the
# old routes remain available as a fallback during migration.


def _org(request):
    profile = getattr(request.user, "profile", None)
    organization = getattr(profile, "organization", None)
    if organization is not None:
        return organization
    organization = m.Organization.objects.first()
    if organization is None:
        organization = m.Organization.objects.create(name="KAYI Haustechnik")
    return organization


def _employee(request, org=None):
    org = org or _org(request)
    employee = getattr(request.user, "employee", None)
    if employee is not None and employee.organization_id == org.id:
        return employee
    email = getattr(request.user, "email", "") or ""
    if email:
        employee = m.Employee.objects.filter(organization=org, email__iexact=email).first()
    return employee


def _role(request):
    profile = getattr(request.user, "profile", None)
    return getattr(profile, "role", "office") or "office"


def _is_field_user(request):
    return _role(request) == "technician" or bool(getattr(getattr(request.user, "profile", None), "is_mobile_worker", False))


def _unique_number(model, org, prefix):
    year = timezone.localdate().year
    base = f"{prefix}-{year}-"
    latest = (
        model.objects.filter(organization=org, number__startswith=base)
        .order_by("-number")
        .values_list("number", flat=True)
        .first()
    )
    value = 1
    if latest:
        match = re.search(r"(\d+)$", latest)
        if match:
            value = int(match.group(1)) + 1
    while True:
        number = f"{base}{value:04d}"
        if not model.objects.filter(organization=org, number=number).exists():
            return number
        value += 1


def _money(value):
    try:
        return Decimal(str(value or "0").replace(",", "."))
    except (InvalidOperation, ValueError):
        return Decimal("0")


def _project_total(project):
    invoices = m.Invoice.objects.filter(project=project).exclude(status="cancelled")
    total = Decimal("0")
    for invoice in invoices.prefetch_related("items"):
        for item in invoice.items.all():
            total += item.quantity * item.unit_price * (Decimal("1") + item.tax_rate / Decimal("100"))
    return total


def _quote_total(quote):
    subtotal = sum((item.quantity * item.unit_price for item in quote.items.all()), Decimal("0"))
    discount = subtotal * quote.discount_percent / Decimal("100")
    net = subtotal - discount
    tax = sum((item.quantity * item.unit_price * item.tax_rate / Decimal("100") for item in quote.items.all()), Decimal("0"))
    return {"net": net, "tax": tax, "gross": net + tax}


def _invoice_total(invoice):
    net = sum((item.quantity * item.unit_price for item in invoice.items.all()), Decimal("0"))
    tax = sum((item.quantity * item.unit_price * item.tax_rate / Decimal("100") for item in invoice.items.all()), Decimal("0"))
    paid = invoice.payments.aggregate(total=Sum("amount"))["total"] or Decimal("0")
    return {"net": net, "tax": tax, "gross": net + tax, "paid": paid, "open": max(Decimal("0"), net + tax - paid)}


class StyledModelForm(forms.ModelForm):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        for field in self.fields.values():
            css = "next-control"
            existing = field.widget.attrs.get("class", "")
            field.widget.attrs["class"] = f"{existing} {css}".strip()
            if isinstance(field.widget, forms.Textarea):
                field.widget.attrs.setdefault("rows", 4)


class CustomerForm(StyledModelForm):
    class Meta:
        model = m.Customer
        fields = [
            "type", "company", "salutation", "first_name", "last_name", "email", "phone", "mobile",
            "street", "postal_code", "city", "country", "vat_id", "notes",
        ]
        widgets = {"notes": forms.Textarea(attrs={"rows": 3})}


class ObjectLocationForm(StyledModelForm):
    class Meta:
        model = m.ObjectLocation
        fields = ["name", "street", "postal_code", "city", "floor", "access_notes"]
        widgets = {"access_notes": forms.Textarea(attrs={"rows": 2})}


class ProjectForm(StyledModelForm):
    class Meta:
        model = m.Project
        fields = ["title", "customer", "object_location", "description", "priority", "manager", "members"]
        widgets = {"description": forms.Textarea(attrs={"rows": 3}), "members": forms.SelectMultiple(attrs={"size": 5})}

    def __init__(self, *args, organization=None, **kwargs):
        super().__init__(*args, **kwargs)
        if organization:
            self.fields["customer"].queryset = m.Customer.objects.filter(organization=organization, active=True)
            self.fields["object_location"].queryset = m.ObjectLocation.objects.filter(organization=organization)
            self.fields["manager"].queryset = m.Employee.objects.filter(organization=organization, active=True)
            self.fields["members"].queryset = m.Employee.objects.filter(organization=organization, active=True)


class AppointmentForm(StyledModelForm):
    class Meta:
        model = m.CalendarEvent
        fields = ["title", "type", "starts_at", "ends_at", "all_day", "location", "notes", "project", "attendees"]
        widgets = {
            "starts_at": forms.DateTimeInput(attrs={"type": "datetime-local"}, format="%Y-%m-%dT%H:%M"),
            "ends_at": forms.DateTimeInput(attrs={"type": "datetime-local"}, format="%Y-%m-%dT%H:%M"),
            "notes": forms.Textarea(attrs={"rows": 3}),
            "attendees": forms.SelectMultiple(attrs={"size": 5}),
        }

    def __init__(self, *args, organization=None, **kwargs):
        super().__init__(*args, **kwargs)
        self.fields["starts_at"].input_formats = ["%Y-%m-%dT%H:%M"]
        self.fields["ends_at"].input_formats = ["%Y-%m-%dT%H:%M"]
        if organization:
            self.fields["project"].queryset = m.Project.objects.filter(organization=organization, archived=False)
            self.fields["attendees"].queryset = m.Employee.objects.filter(organization=organization, active=True)


class QuoteForm(StyledModelForm):
    class Meta:
        model = m.Quote
        fields = ["project", "issue_date", "valid_until", "intro_text", "outro_text", "discount_percent", "notes"]
        widgets = {
            "issue_date": forms.DateInput(attrs={"type": "date"}),
            "valid_until": forms.DateInput(attrs={"type": "date"}),
            "intro_text": forms.Textarea(attrs={"rows": 2}),
            "outro_text": forms.Textarea(attrs={"rows": 2}),
            "notes": forms.Textarea(attrs={"rows": 2}),
        }

    def __init__(self, *args, organization=None, **kwargs):
        super().__init__(*args, **kwargs)
        if organization:
            self.fields["project"].queryset = m.Project.objects.filter(organization=organization, archived=False)


class InvoiceForm(StyledModelForm):
    class Meta:
        model = m.Invoice
        fields = ["project", "quote", "issue_date", "due_date", "service_date", "intro_text", "outro_text", "notes"]
        widgets = {
            "issue_date": forms.DateInput(attrs={"type": "date"}),
            "due_date": forms.DateInput(attrs={"type": "date"}),
            "service_date": forms.DateInput(attrs={"type": "date"}),
            "intro_text": forms.Textarea(attrs={"rows": 2}),
            "outro_text": forms.Textarea(attrs={"rows": 2}),
            "notes": forms.Textarea(attrs={"rows": 2}),
        }

    def __init__(self, *args, organization=None, **kwargs):
        super().__init__(*args, **kwargs)
        if organization:
            self.fields["project"].queryset = m.Project.objects.filter(organization=organization, archived=False)
            self.fields["quote"].queryset = m.Quote.objects.filter(organization=organization)


@login_required
def dashboard(request):
    if _is_field_user(request):
        return redirect("next-field")
    org = _org(request)
    today = timezone.localdate()
    start = timezone.make_aware(timezone.datetime.combine(today, timezone.datetime.min.time()))
    end = start + timedelta(days=1)
    appointments = (
        m.CalendarEvent.objects.filter(organization=org, starts_at__gte=start, starts_at__lt=end)
        .select_related("project", "project__customer")
        .prefetch_related("attendees")
        .order_by("starts_at")[:8]
    )
    open_quotes = m.Quote.objects.filter(organization=org, status__in=["draft", "review", "sent"]).count()
    overdue = m.Invoice.objects.filter(organization=org, status="overdue").count()
    active_projects = m.Project.objects.filter(organization=org, archived=False).exclude(status__in=["completed", "cancelled"]).count()
    recent_projects = m.Project.objects.filter(organization=org, archived=False).select_related("customer").order_by("-updated_at")[:6]
    return render(request, "rebuild/dashboard.html", {
        "appointments": appointments,
        "open_quotes": open_quotes,
        "overdue": overdue,
        "active_projects": active_projects,
        "recent_projects": recent_projects,
        "today": today,
    })


@login_required
def customer_list(request):
    org = _org(request)
    query = request.GET.get("q", "").strip()
    customers = m.Customer.objects.filter(organization=org, active=True)
    if query:
        customers = customers.filter(
            Q(company__icontains=query) | Q(first_name__icontains=query) | Q(last_name__icontains=query)
            | Q(email__icontains=query) | Q(phone__icontains=query) | Q(number__icontains=query)
        )
    customers = customers.order_by("company", "last_name", "first_name")[:250]
    return render(request, "rebuild/customers.html", {"customers": customers, "query": query})


@login_required
@require_http_methods(["GET", "POST"])
def customer_create(request):
    org = _org(request)
    form = CustomerForm(request.POST or None)
    location_form = ObjectLocationForm(request.POST or None, prefix="site")
    if request.method == "POST" and form.is_valid():
        with transaction.atomic():
            customer = form.save(commit=False)
            customer.organization = org
            customer.number = _unique_number(m.Customer, org, "K")
            customer.save()
            if request.POST.get("site-street") and location_form.is_valid():
                location = location_form.save(commit=False)
                location.organization = org
                location.customer = customer
                location.save()
        messages.success(request, "Kunde wurde angelegt.")
        return redirect("next-customer-detail", pk=customer.pk)
    return render(request, "rebuild/customer_form.html", {"form": form, "location_form": location_form, "mode": "create"})


@login_required
@require_http_methods(["GET", "POST"])
def customer_detail(request, pk):
    org = _org(request)
    customer = get_object_or_404(m.Customer, pk=pk, organization=org)
    if request.method == "POST":
        form = CustomerForm(request.POST, instance=customer)
        if form.is_valid():
            form.save()
            messages.success(request, "Kundendaten gespeichert.")
            return redirect("next-customer-detail", pk=customer.pk)
    else:
        form = CustomerForm(instance=customer)
    projects = customer.projects.filter(organization=org).order_by("-updated_at")
    locations = customer.object_locations.all()
    return render(request, "rebuild/customer_detail.html", {"customer": customer, "form": form, "projects": projects, "locations": locations})


@login_required
def project_list(request):
    org = _org(request)
    query = request.GET.get("q", "").strip()
    status = request.GET.get("status", "").strip()
    projects = m.Project.objects.filter(organization=org, archived=False).select_related("customer", "manager", "object_location")
    if query:
        projects = projects.filter(Q(number__icontains=query) | Q(title__icontains=query) | Q(customer__company__icontains=query) | Q(customer__last_name__icontains=query))
    if status:
        projects = projects.filter(status=status)
    return render(request, "rebuild/projects.html", {"projects": projects.order_by("-updated_at")[:250], "query": query, "status": status, "statuses": m.Project._meta.get_field("status").choices})


@login_required
@require_http_methods(["GET", "POST"])
def project_create(request):
    org = _org(request)
    initial = {}
    if request.GET.get("customer"):
        initial["customer"] = request.GET.get("customer")
    form = ProjectForm(request.POST or None, organization=org, initial=initial)
    if request.method == "POST" and form.is_valid():
        project = form.save(commit=False)
        project.organization = org
        project.number = _unique_number(m.Project, org, "P")
        project.status = "inquiry"
        project.save()
        form.save_m2m()
        messages.success(request, "Projekt angelegt. Du kannst jetzt Termin, Angebot oder Dokumentation hinzufügen.")
        return redirect("next-project-detail", pk=project.pk)
    return render(request, "rebuild/project_form.html", {"form": form})


@login_required
def project_detail(request, pk):
    org = _org(request)
    project = get_object_or_404(m.Project.objects.select_related("customer", "object_location", "manager"), pk=pk, organization=org)
    appointments = project.events.prefetch_related("attendees").order_by("-starts_at")[:20]
    quotes = project.quotes.prefetch_related("items").order_by("-created_at")
    invoices = project.invoices.prefetch_related("items", "payments").order_by("-created_at")
    documents = project.documents.order_by("-created_at")[:30]
    tasks = project.tasks.order_by("status", "due_at")[:20]
    materials = project.materials.order_by("-created_at")[:20]
    invoice_gross = sum((_invoice_total(invoice)["gross"] for invoice in invoices), Decimal("0"))
    return render(request, "rebuild/project_detail.html", {
        "project": project, "appointments": appointments, "quotes": quotes, "invoices": invoices,
        "documents": documents, "tasks": tasks, "materials": materials, "invoice_gross": invoice_gross,
    })


@login_required
def appointment_list(request):
    org = _org(request)
    start = request.GET.get("start")
    events = m.CalendarEvent.objects.filter(organization=org).select_related("project", "project__customer").prefetch_related("attendees")
    if start:
        try:
            start_date = timezone.datetime.strptime(start, "%Y-%m-%d").date()
        except ValueError:
            start_date = timezone.localdate()
    else:
        start_date = timezone.localdate()
    week_start = start_date - timedelta(days=start_date.weekday())
    week_end = week_start + timedelta(days=7)
    aware_start = timezone.make_aware(timezone.datetime.combine(week_start, timezone.datetime.min.time()))
    aware_end = timezone.make_aware(timezone.datetime.combine(week_end, timezone.datetime.min.time()))
    events = events.filter(starts_at__gte=aware_start, starts_at__lt=aware_end).order_by("starts_at")
    days = []
    for offset in range(7):
        date = week_start + timedelta(days=offset)
        days.append({"date": date, "events": [event for event in events if timezone.localtime(event.starts_at).date() == date]})
    return render(request, "rebuild/appointments.html", {"days": days, "week_start": week_start, "prev": week_start - timedelta(days=7), "next": week_start + timedelta(days=7)})


@login_required
@require_http_methods(["GET", "POST"])
def appointment_create(request):
    org = _org(request)
    now = timezone.localtime().replace(second=0, microsecond=0)
    initial = {"starts_at": now + timedelta(hours=1), "ends_at": now + timedelta(hours=2)}
    if request.GET.get("project"):
        initial["project"] = request.GET.get("project")
    form = AppointmentForm(request.POST or None, organization=org, initial=initial)
    if request.method == "POST" and form.is_valid():
        event = form.save(commit=False)
        event.organization = org
        event.created_by = request.user
        event.save()
        form.save_m2m()
        messages.success(request, "Termin wurde geplant.")
        return redirect("next-appointment-detail", pk=event.pk)
    return render(request, "rebuild/appointment_form.html", {"form": form})


@login_required
def appointment_detail(request, pk):
    org = _org(request)
    event = get_object_or_404(m.CalendarEvent.objects.select_related("project", "project__customer", "project__object_location"), pk=pk, organization=org)
    docs = m.Document.objects.filter(organization=org, project=event.project, metadata__event_id=event.pk).order_by("-created_at") if event.project else m.Document.objects.none()
    employee = _employee(request, org)
    running = None
    if employee and event.project:
        running = m.TimeEntry.objects.filter(organization=org, employee=employee, project=event.project, ended_at__isnull=True).order_by("-started_at").first()
    return render(request, "rebuild/appointment_detail.html", {"event": event, "documents": docs, "running": running, "employee": employee})


@login_required
def field_home(request):
    org = _org(request)
    employee = _employee(request, org)
    events = m.CalendarEvent.objects.filter(organization=org).select_related("project", "project__customer").prefetch_related("attendees")
    if employee:
        events = events.filter(attendees=employee)
    now = timezone.now()
    planned = events.filter(ends_at__gte=now).order_by("starts_at")[:30]
    overdue = events.filter(ends_at__lt=now).order_by("-starts_at")[:30]
    documented_ids = set(
        m.Document.objects.filter(organization=org, category="report", metadata__event_id__isnull=False)
        .values_list("metadata__event_id", flat=True)
    )
    overdue = [event for event in overdue if event.pk not in documented_ids]
    documented = [event for event in events.order_by("-starts_at")[:100] if event.pk in documented_ids][:30]
    return render(request, "rebuild/field_home.html", {"planned": planned, "overdue": overdue, "documented": documented, "employee": employee})


@login_required
@require_POST
def time_toggle(request, event_pk):
    org = _org(request)
    event = get_object_or_404(m.CalendarEvent, pk=event_pk, organization=org)
    employee = _employee(request, org)
    if employee is None or event.project_id is None:
        return JsonResponse({"ok": False, "error": "Mitarbeiter oder Projekt fehlt."}, status=400)
    running = m.TimeEntry.objects.filter(organization=org, employee=employee, project=event.project, ended_at__isnull=True).order_by("-started_at").first()
    if running:
        running.ended_at = timezone.now()
        running.save(update_fields=["ended_at", "updated_at"])
        return JsonResponse({"ok": True, "state": "stopped"})
    entry = m.TimeEntry.objects.create(
        organization=org, employee=employee, project=event.project, started_at=timezone.now(),
        description=f"Termin #{event.pk}: {event.title}",
    )
    return JsonResponse({"ok": True, "state": "running", "id": entry.pk})


@login_required
@require_POST
def appointment_document(request, pk):
    org = _org(request)
    event = get_object_or_404(m.CalendarEvent.objects.select_related("project", "project__customer"), pk=pk, organization=org)
    if event.project_id is None:
        return JsonResponse({"ok": False, "error": "Der Termin muss einem Projekt zugeordnet sein."}, status=400)
    report_text = (request.POST.get("report_text") or "").strip()
    services = (request.POST.get("services") or "").strip()
    material = (request.POST.get("material") or "").strip()
    customer_name = (request.POST.get("customer_name") or "").strip()
    payload = {
        "event_id": event.pk,
        "event_title": event.title,
        "services": services,
        "material": material,
        "customer_name": customer_name,
        "source": "kayi-next-field",
    }
    body = report_text or "Vor-Ort-Dokumentation"
    document = m.Document(
        organization=org,
        customer=event.project.customer,
        project=event.project,
        title=f"Arbeitsbericht · {event.title} · {timezone.localdate():%d.%m.%Y}",
        category="report",
        mime_type="text/plain",
        size=len(body.encode("utf-8")),
        metadata=payload,
        uploaded_by=request.user,
    )
    document.file.save(f"arbeitsbericht-{event.pk}-{timezone.now():%Y%m%d%H%M%S}.txt", ContentFile(body.encode("utf-8")), save=False)
    document.save()

    for upload in request.FILES.getlist("photos"):
        photo = m.Document(
            organization=org, customer=event.project.customer, project=event.project,
            title=upload.name, category="photo", mime_type=getattr(upload, "content_type", "") or "",
            size=getattr(upload, "size", 0) or 0, metadata={"event_id": event.pk, "source": "kayi-next-field"}, uploaded_by=request.user,
        )
        photo.file.save(upload.name, upload, save=False)
        photo.save()

    signature_data = request.POST.get("signature_data") or ""
    if signature_data.startswith("data:image/png;base64,"):
        try:
            raw = base64.b64decode(signature_data.split(",", 1)[1])
            signature = m.Document(
                organization=org, customer=event.project.customer, project=event.project,
                title=f"Kundenunterschrift · {customer_name or event.project.customer.display_name}", category="other",
                mime_type="image/png", size=len(raw), metadata={"event_id": event.pk, "kind": "customer_signature"}, uploaded_by=request.user,
            )
            signature.file.save(f"signature-{event.pk}.png", ContentFile(raw), save=False)
            signature.save()
        except Exception:
            pass

    if event.project.status in {"inquiry", "planning", "quoted", "confirmed"}:
        event.project.status = "in_progress"
        event.project.actual_start = event.project.actual_start or timezone.localdate()
        event.project.save(update_fields=["status", "actual_start", "updated_at"])
    return JsonResponse({"ok": True, "redirect": f"/appointments/{event.pk}/"})


@login_required
@require_POST
def ai_structure_report(request, pk):
    org = _org(request)
    get_object_or_404(m.CalendarEvent, pk=pk, organization=org)
    raw = (request.POST.get("text") or "").strip()
    if not raw:
        return JsonResponse({"ok": False, "error": "Kein Text vorhanden."}, status=400)
    fallback = {"report": raw, "services": "", "material": ""}
    key = os.environ.get("OPENAI_API_KEY", "").strip()
    if not key:
        return JsonResponse({"ok": True, **fallback, "ai": False})
    try:
        from openai import OpenAI
        client = OpenAI(api_key=key)
        response = client.responses.create(
            model=os.environ.get("OPENAI_MODEL", "gpt-4.1-mini"),
            input=[
                {"role": "system", "content": "Du strukturierst Baustellen-Diktate für einen deutschen Haustechnikbetrieb. Antworte ausschließlich als JSON mit report, services, material. Erfinde nichts."},
                {"role": "user", "content": raw},
            ],
        )
        text = response.output_text.strip()
        match = re.search(r"\{.*\}", text, re.S)
        data = json.loads(match.group(0) if match else text)
        return JsonResponse({"ok": True, "report": data.get("report", raw), "services": data.get("services", ""), "material": data.get("material", ""), "ai": True})
    except Exception:
        return JsonResponse({"ok": True, **fallback, "ai": False})


@login_required
def quote_list(request):
    org = _org(request)
    quotes = m.Quote.objects.filter(organization=org).select_related("project", "project__customer").prefetch_related("items").order_by("-created_at")[:250]
    rows = [{"quote": quote, "total": _quote_total(quote)} for quote in quotes]
    return render(request, "rebuild/quotes.html", {"rows": rows})


def _save_quote_items(quote, request):
    descriptions = request.POST.getlist("item_description")
    quantities = request.POST.getlist("item_quantity")
    units = request.POST.getlist("item_unit")
    prices = request.POST.getlist("item_price")
    taxes = request.POST.getlist("item_tax")
    quote.items.all().delete()
    position = 1
    for index, description in enumerate(descriptions):
        description = description.strip()
        if not description:
            continue
        m.QuoteItem.objects.create(
            quote=quote, position=position, description=description,
            quantity=_money(quantities[index] if index < len(quantities) else 1),
            unit=(units[index] if index < len(units) else "Stk.") or "Stk.",
            unit_price=_money(prices[index] if index < len(prices) else 0),
            tax_rate=_money(taxes[index] if index < len(taxes) else 19),
        )
        position += 1


@login_required
@require_http_methods(["GET", "POST"])
def quote_editor(request, pk=None):
    org = _org(request)
    quote = get_object_or_404(m.Quote, pk=pk, organization=org) if pk else None
    initial = {}
    if request.GET.get("project"):
        initial["project"] = request.GET.get("project")
    form = QuoteForm(request.POST or None, instance=quote, organization=org, initial=initial)
    if request.method == "POST" and form.is_valid():
        with transaction.atomic():
            obj = form.save(commit=False)
            obj.organization = org
            obj.created_by = obj.created_by or request.user
            if not obj.number:
                obj.number = _unique_number(m.Quote, org, "A")
            if request.POST.get("action") == "send":
                obj.status = "sent"
                obj.sent_at = timezone.now()
            obj.save()
            _save_quote_items(obj, request)
        messages.success(request, "Angebot gespeichert.")
        return redirect("next-quote-edit", pk=obj.pk)
    catalog = m.CatalogItem.objects.filter(organization=org, active=True).order_by("name")[:500]
    return render(request, "rebuild/document_editor.html", {"form": form, "document": quote, "items": quote.items.all() if quote else [], "catalog": catalog, "kind": "quote", "totals": _quote_total(quote) if quote else None})


@login_required
def invoice_list(request):
    org = _org(request)
    invoices = m.Invoice.objects.filter(organization=org).select_related("project", "project__customer").prefetch_related("items", "payments").order_by("-created_at")[:250]
    rows = [{"invoice": invoice, "total": _invoice_total(invoice)} for invoice in invoices]
    return render(request, "rebuild/invoices.html", {"rows": rows})


def _save_invoice_items(invoice, request):
    descriptions = request.POST.getlist("item_description")
    quantities = request.POST.getlist("item_quantity")
    units = request.POST.getlist("item_unit")
    prices = request.POST.getlist("item_price")
    taxes = request.POST.getlist("item_tax")
    invoice.items.all().delete()
    position = 1
    for index, description in enumerate(descriptions):
        description = description.strip()
        if not description:
            continue
        m.InvoiceItem.objects.create(
            invoice=invoice, position=position, description=description,
            quantity=_money(quantities[index] if index < len(quantities) else 1),
            unit=(units[index] if index < len(units) else "Stk.") or "Stk.",
            unit_price=_money(prices[index] if index < len(prices) else 0),
            tax_rate=_money(taxes[index] if index < len(taxes) else 19),
        )
        position += 1


@login_required
@require_http_methods(["GET", "POST"])
def invoice_editor(request, pk=None):
    org = _org(request)
    invoice = get_object_or_404(m.Invoice, pk=pk, organization=org) if pk else None
    initial = {"due_date": timezone.localdate() + timedelta(days=14)}
    if request.GET.get("project"):
        initial["project"] = request.GET.get("project")
    if request.GET.get("quote"):
        initial["quote"] = request.GET.get("quote")
    form = InvoiceForm(request.POST or None, instance=invoice, organization=org, initial=initial)
    if request.method == "POST" and form.is_valid():
        with transaction.atomic():
            obj = form.save(commit=False)
            obj.organization = org
            obj.created_by = obj.created_by or request.user
            if not obj.number:
                obj.number = _unique_number(m.Invoice, org, "R")
            if request.POST.get("action") == "send":
                obj.status = "sent"
                obj.sent_at = timezone.now()
            obj.save()
            _save_invoice_items(obj, request)
        messages.success(request, "Rechnung gespeichert.")
        return redirect("next-invoice-edit", pk=obj.pk)
    return render(request, "rebuild/document_editor.html", {"form": form, "document": invoice, "items": invoice.items.all() if invoice else [], "catalog": m.CatalogItem.objects.filter(organization=org, active=True).order_by("name")[:500], "kind": "invoice", "totals": _invoice_total(invoice) if invoice else None})


@login_required
@require_POST
def invoice_payment(request, pk):
    org = _org(request)
    invoice = get_object_or_404(m.Invoice, pk=pk, organization=org)
    amount = _money(request.POST.get("amount"))
    if amount <= 0:
        messages.error(request, "Bitte einen gültigen Betrag eingeben.")
        return redirect("next-invoice-edit", pk=invoice.pk)
    m.Payment.objects.create(invoice=invoice, amount=amount, paid_at=timezone.localdate(), method=request.POST.get("method") or "Überweisung", reference=request.POST.get("reference") or "", recorded_by=request.user)
    totals = _invoice_total(invoice)
    invoice.status = "paid" if totals["open"] <= 0 else "partial"
    invoice.save(update_fields=["status", "updated_at"])
    messages.success(request, "Zahlung erfasst.")
    return redirect("next-invoice-edit", pk=invoice.pk)


@login_required
def time_overview(request):
    org = _org(request)
    employee = _employee(request, org)
    entries = m.TimeEntry.objects.filter(organization=org).select_related("employee", "project")
    if _is_field_user(request) and employee:
        entries = entries.filter(employee=employee)
    entries = entries.order_by("-started_at")[:150]
    return render(request, "rebuild/time_overview.html", {"entries": entries, "employee": employee})


@login_required
def migration_import(request):
    org = _org(request)
    summary = None
    errors = []
    if request.method == "POST" and request.FILES.get("file"):
        file = request.FILES["file"]
        kind = request.POST.get("kind") or "customers"
        try:
            rows = _read_table(file)
            summary = _import_tooltime_rows(org, request.user, kind, rows)
            messages.success(request, f"ToolTime-Import abgeschlossen: {summary['created']} neu, {summary['updated']} aktualisiert.")
        except Exception as exc:
            errors.append(str(exc))
    return render(request, "rebuild/migration.html", {"summary": summary, "errors": errors})


def _read_table(upload):
    name = upload.name.lower()
    raw = upload.read()
    if name.endswith(".xlsx"):
        try:
            import openpyxl
        except ImportError as exc:
            raise ValueError("XLSX-Unterstützung fehlt. Bitte CSV exportieren oder openpyxl installieren.") from exc
        book = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        sheet = book.active
        values = list(sheet.iter_rows(values_only=True))
        if not values:
            return []
        headers = [str(value or "").strip() for value in values[0]]
        return [dict(zip(headers, row)) for row in values[1:] if any(value not in (None, "") for value in row)]
    text = raw.decode("utf-8-sig", errors="replace")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t,")
    except csv.Error:
        dialect = csv.excel
        dialect.delimiter = ";"
    return list(csv.DictReader(io.StringIO(text), dialect=dialect))


def _norm(row):
    return {re.sub(r"[^a-z0-9]", "", str(key).lower().replace("ä", "ae").replace("ö", "oe").replace("ü", "ue").replace("ß", "ss")): ("" if value is None else str(value).strip()) for key, value in row.items()}


def _pick(row, *aliases):
    for alias in aliases:
        key = re.sub(r"[^a-z0-9]", "", alias.lower().replace("ä", "ae").replace("ö", "oe").replace("ü", "ue").replace("ß", "ss"))
        if row.get(key):
            return row[key]
    return ""


def _import_tooltime_rows(org, user, kind, rows):
    created = updated = skipped = 0
    with transaction.atomic():
        if kind == "customers":
            for source in rows:
                row = _norm(source)
                number = _pick(row, "Kundennummer", "Customer Number", "Nummer")
                email = _pick(row, "E-Mail", "Email")
                company = _pick(row, "Firma", "Unternehmen", "Company")
                first = _pick(row, "Vorname", "First Name")
                last = _pick(row, "Nachname", "Last Name")
                query = m.Customer.objects.filter(organization=org)
                customer = query.filter(number=number).first() if number else None
                if customer is None and email:
                    customer = query.filter(email__iexact=email).first()
                values = {
                    "company": company, "first_name": first, "last_name": last, "email": email,
                    "phone": _pick(row, "Telefon", "Phone"), "mobile": _pick(row, "Mobil", "Mobile"),
                    "street": _pick(row, "Straße", "Strasse", "Street"), "postal_code": _pick(row, "PLZ", "Postleitzahl", "ZIP"),
                    "city": _pick(row, "Ort", "Stadt", "City"), "country": _pick(row, "Land", "Country") or "DE",
                }
                if customer:
                    for key, value in values.items():
                        if value:
                            setattr(customer, key, value)
                    customer.save()
                    updated += 1
                else:
                    m.Customer.objects.create(organization=org, number=number or _unique_number(m.Customer, org, "K"), **values)
                    created += 1
        elif kind in {"quotes", "invoices"}:
            model = m.Quote if kind == "quotes" else m.Invoice
            prefix = "A" if kind == "quotes" else "R"
            grouped = {}
            for source in rows:
                row = _norm(source)
                number = _pick(row, "Angebotsnummer" if kind == "quotes" else "Rechnungsnummer", "Nummer", "Number")
                if not number:
                    skipped += 1
                    continue
                grouped.setdefault(number, []).append(row)
            for number, group in grouped.items():
                first = group[0]
                customer_number = _pick(first, "Kundennummer", "Customer Number")
                customer_name = _pick(first, "Kunde", "Kundenname", "Customer")
                customer = m.Customer.objects.filter(organization=org, number=customer_number).first() if customer_number else None
                if customer is None and customer_name:
                    customer = m.Customer.objects.filter(organization=org).filter(Q(company__iexact=customer_name) | Q(last_name__iexact=customer_name)).first()
                if customer is None:
                    customer = m.Customer.objects.create(organization=org, number=_unique_number(m.Customer, org, "K"), company=customer_name or "ToolTime Import")
                project_title = _pick(first, "Projekt", "Projekttitel", "Project") or f"ToolTime Import {number}"
                project = m.Project.objects.filter(organization=org, customer=customer, title=project_title).first()
                if project is None:
                    project = m.Project.objects.create(organization=org, customer=customer, number=_unique_number(m.Project, org, "P"), title=project_title, status="invoiced" if kind == "invoices" else "quoted")
                obj = model.objects.filter(organization=org, number=number).first()
                if obj:
                    updated += 1
                else:
                    kwargs = {"organization": org, "number": number or _unique_number(model, org, prefix), "project": project, "created_by": user}
                    if kind == "invoices":
                        kwargs["due_date"] = timezone.localdate() + timedelta(days=14)
                    obj = model.objects.create(**kwargs)
                    created += 1
                item_model = m.QuoteItem if kind == "quotes" else m.InvoiceItem
                parent_field = "quote" if kind == "quotes" else "invoice"
                if not getattr(obj, "items").exists():
                    for position, row in enumerate(group, 1):
                        description = _pick(row, "Positionsbezeichnung", "Beschreibung", "Description", "Leistung", "Artikel")
                        if not description:
                            continue
                        item_model.objects.create(**{
                            parent_field: obj, "position": position, "description": description,
                            "quantity": _money(_pick(row, "Menge", "Quantity") or 1),
                            "unit": _pick(row, "Einheit", "Unit") or "Stk.",
                            "unit_price": _money(_pick(row, "Einzelpreis", "Preis", "Unit Price", "Netto")),
                            "tax_rate": _money(_pick(row, "MwSt", "Steuer", "Tax") or 19),
                        })
        elif kind == "time":
            employee = m.Employee.objects.filter(organization=org, user=user).first() or m.Employee.objects.filter(organization=org).first()
            if employee is None:
                raise ValueError("Für den Zeitimport muss mindestens ein Mitarbeiter existieren.")
            for source in rows:
                row = _norm(source)
                project_number = _pick(row, "Projektnummer", "Project Number")
                project = m.Project.objects.filter(organization=org, number=project_number).first() if project_number else None
                if project is None:
                    skipped += 1
                    continue
                date_text = _pick(row, "Datum", "Date")
                start_text = _pick(row, "Start", "Von", "Startzeit")
                end_text = _pick(row, "Ende", "Bis", "Endzeit")
                try:
                    date = timezone.datetime.strptime(date_text, "%d.%m.%Y").date() if "." in date_text else timezone.datetime.fromisoformat(date_text).date()
                    start_time = timezone.datetime.strptime(start_text, "%H:%M").time()
                    end_time = timezone.datetime.strptime(end_text, "%H:%M").time()
                except Exception:
                    skipped += 1
                    continue
                started = timezone.make_aware(timezone.datetime.combine(date, start_time))
                ended = timezone.make_aware(timezone.datetime.combine(date, end_time))
                m.TimeEntry.objects.create(organization=org, employee=employee, project=project, started_at=started, ended_at=ended, break_minutes=int(_money(_pick(row, "Pause", "Pausenminuten") or 0)), description=_pick(row, "Beschreibung", "Description"))
                created += 1
        else:
            raise ValueError("Unbekannter Importtyp.")
    return {"created": created, "updated": updated, "skipped": skipped, "rows": len(rows)}


@login_required
def settings_page(request):
    org = _org(request)
    integrations = m.IntegrationConfig.objects.filter(organization=org).order_by("provider")
    return render(request, "rebuild/settings.html", {"organization": org, "integrations": integrations})
