from __future__ import annotations

import csv
import hashlib
import mimetypes
import re
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterable

import fitz
from django.core.files import File
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from erp.models import CatalogItem, Organization, PriceItem, PriceSource, Supplier


PRICE_EXTENSIONS = {".xlsx", ".xlsm", ".csv", ".pdf", ".003", ".p86"}
CODE_NAMES = {"code", "nr", "nummer", "pos", "position", "positionsnummer", "artikelnummer", "leistungsnummer", "lvpos"}
DESCRIPTION_NAMES = {"beschreibung", "bezeichnung", "leistung", "text", "langtext", "artikel", "name", "kurztext"}
UNIT_NAMES = {"einheit", "me", "mengeneinheit", "unit", "eh"}
PURCHASE_NAMES = {"ek", "einkauf", "einkaufspreis", "nettoeinkauf", "preisnetto", "nettoek"}
SALES_NAMES = {"vk", "verkauf", "verkaufspreis", "preis", "einheitspreis", "ep", "netto", "nettopreis", "listenpreis_netto", "listenpreisnetto", "neu"}
CATEGORY_NAMES = {"kategorie", "gewerk", "gruppe", "warengruppe", "kapitel", "bereich"}


def _norm(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").strip().lower().replace("ä", "ae").replace("ö", "oe").replace("ü", "ue").replace("ß", "ss"))


def _decimal(value: object) -> Decimal | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value)).quantize(Decimal("0.01"))
    text = str(value).strip().replace("€", "").replace("EUR", "").replace(" ", "")
    if not text:
        return None
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        text = text.replace(".", "").replace(",", ".")
    text = re.sub(r"[^0-9.\-]", "", text)
    try:
        return Decimal(text).quantize(Decimal("0.01")) if text not in {"", "-", "."} else None
    except InvalidOperation:
        return None


def _pick(headers: list[str], names: set[str]) -> int | None:
    normalized = [_norm(h) for h in headers]
    targets = {_norm(x) for x in names}
    for idx, header in enumerate(normalized):
        if header in targets:
            return idx
    for idx, header in enumerate(normalized):
        if any(target and target in header for target in targets):
            return idx
    return None


def _rows_to_items(rows: list[list[object]], sheet: str = "") -> list[dict]:
    if not rows:
        return []
    header_idx = 0
    best_score = -1
    for idx, row in enumerate(rows[:30]):
        vals = [_norm(v) for v in row]
        score = sum(any(n in vals for n in {_norm(x) for x in group}) for group in (CODE_NAMES, DESCRIPTION_NAMES, UNIT_NAMES, SALES_NAMES))
        if score > best_score:
            best_score, header_idx = score, idx
    headers = [str(v or "") for v in rows[header_idx]]
    code_i = _pick(headers, CODE_NAMES)
    desc_i = _pick(headers, DESCRIPTION_NAMES)
    unit_i = _pick(headers, UNIT_NAMES)
    purchase_i = _pick(headers, PURCHASE_NAMES)
    sales_i = _pick(headers, SALES_NAMES)
    category_i = _pick(headers, CATEGORY_NAMES)
    items: list[dict] = []
    for raw in rows[header_idx + 1 :]:
        values = list(raw) + [None] * max(0, len(headers) - len(raw))
        code = str(values[code_i] or "").strip() if code_i is not None else ""
        description = str(values[desc_i] or "").strip() if desc_i is not None else ""
        if not description:
            nonempty = [str(v).strip() for v in values if v not in (None, "")]
            if len(nonempty) >= 2:
                description = nonempty[1] if code and nonempty[0] == code else nonempty[0]
        if not description or _norm(description) in {_norm(h) for h in headers}:
            continue
        purchase = _decimal(values[purchase_i]) if purchase_i is not None else None
        sales = _decimal(values[sales_i]) if sales_i is not None else None
        if not code and purchase is None and sales is None and len(description) < 4:
            continue
        items.append({
            "code": code[:120],
            "description": description[:5000],
            "category": (str(values[category_i] or "").strip() if category_i is not None else sheet)[:180],
            "unit": (str(values[unit_i] or "").strip() if unit_i is not None else "")[:40],
            "purchase_price": purchase,
            "sales_price": sales,
            "external_data": {"sheet": sheet, "headers": headers[:40]},
        })
    return items


def _xlsx_items(path: Path) -> list[dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    result: list[dict] = []
    for sheet in workbook.worksheets:
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        result.extend(_rows_to_items(rows, sheet.title))
    return result


def _text_rows(path: Path) -> list[list[str]]:
    raw = path.read_bytes()
    text = None
    for encoding in ("utf-8-sig", "cp1252", "latin1"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    text = text or raw.decode("latin1", errors="replace")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ";" if sample.count(";") >= sample.count("\t") else "\t"
    return [row for row in csv.reader(text.splitlines(), delimiter=delimiter)]


def _pdf_items(path: Path) -> list[dict]:
    doc = fitz.open(path)
    lines: list[str] = []
    for page in doc:
        lines.extend(line.strip() for line in page.get_text("text").splitlines() if line.strip())
    result: list[dict] = []
    price_pattern = re.compile(r"(?P<price>\d{1,6}(?:[.,]\d{2}))\s*(?:€|EUR)?\s*$", re.I)
    code_pattern = re.compile(r"^(?P<code>[A-Z0-9][A-Z0-9._/\-]{1,30})\s+", re.I)
    for line in lines:
        match = price_pattern.search(line)
        if not match:
            continue
        code_match = code_pattern.search(line)
        code = code_match.group("code") if code_match else ""
        description = line[: match.start()].strip(" .;:-")
        if code and description.startswith(code):
            description = description[len(code):].strip(" .;:-")
        if description:
            result.append({"code": code, "description": description, "category": path.stem[:180], "unit": "", "purchase_price": None, "sales_price": _decimal(match.group("price")), "external_data": {"source": "pdf"}})
    return result


def _detect(path: Path) -> tuple[str, str, str]:
    text = str(path).lower().replace("\\", "/")
    filename = path.name.lower()
    if "allianz" in text:
        return PriceSource.Kind.INSURANCE, "Allianz", "Allianz"
    # B&O ships operational prices and several analysis/mapping helpers in the
    # same folder. Only PL1658/VA04 and the main workbook are selectable as the
    # commercial insurance basis.
    if any(token in filename for token in ("mapping", "preise_beobachtet", "freiposition")):
        return PriceSource.Kind.MAPPING, path.stem, "B&O"
    if any(token in text for token in ("/bo/", "b&o", "geschaeftskunde", "geschäftskunde")):
        if any(token in filename for token in ("pl1658", "preisliste", "lv preise", "va04")):
            return PriceSource.Kind.INSURANCE, "B&O PL 1658", "B&O"
        return PriceSource.Kind.MAPPING, path.stem, "B&O"
    if "joka" in text:
        return PriceSource.Kind.SUPPLIER, "JOKA", "JOKA"
    if "raab" in text or "karcher" in text:
        return PriceSource.Kind.SUPPLIER, "Raab Karcher", "Raab Karcher"
    if "privat" in text:
        return PriceSource.Kind.CUSTOMER, "Privatkunden", ""
    if "leistungskatalog" in text:
        return PriceSource.Kind.CATALOG, "A+Bau Leistungskatalog", ""
    return PriceSource.Kind.OTHER, path.stem, ""


@transaction.atomic
def import_price_file(path: str | Path, organization: Organization, *, copy_raw: bool = True) -> dict:
    path = Path(path)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    kind, display_name, supplier_name = _detect(path)
    supplier = None
    if supplier_name:
        supplier, _ = Supplier.objects.get_or_create(
            organization=organization,
            name=supplier_name,
            defaults={"number": re.sub(r"[^A-Z0-9]", "", supplier_name.upper())[:30] or f"L{Supplier.objects.filter(organization=organization).count()+1:04d}"},
        )
    source, created = PriceSource.objects.get_or_create(
        organization=organization,
        sha256=digest,
        defaults={
            "supplier": supplier,
            "name": display_name,
            "kind": kind,
            "original_filename": path.name,
            "mime_type": mimetypes.guess_type(path.name)[0] or "application/octet-stream",
        },
    )
    if not created and source.items.exists():
        return {"source": source.name, "created": False, "rows": source.imported_rows, "sha256": digest}
    if copy_raw and (not source.raw_file or not source.raw_file.name):
        with path.open("rb") as handle:
            source.raw_file.save(path.name, File(handle), save=False)
    suffix = path.suffix.lower()
    try:
        if suffix in {".xlsx", ".xlsm"}:
            items = _xlsx_items(path)
        elif suffix in {".csv", ".003", ".p86"}:
            items = _rows_to_items(_text_rows(path), path.stem)
        elif suffix == ".pdf":
            items = _pdf_items(path)
        else:
            items = []
    except Exception as exc:
        items = []
        source.import_summary = {"error": str(exc)}
    source.items.all().delete()
    PriceItem.objects.bulk_create([
        PriceItem(organization=organization, source=source, **item) for item in items
    ], batch_size=1000)
    source.supplier = supplier
    source.name = display_name
    source.kind = kind
    source.imported_at = timezone.now()
    source.imported_rows = len(items)
    source.import_summary = {**(source.import_summary or {}), "extension": suffix, "rows": len(items), "path": str(path.parent.name)}
    source.save()

    if kind == PriceSource.Kind.CATALOG:
        for item in items:
            if not item["code"]:
                continue
            CatalogItem.objects.update_or_create(
                organization=organization,
                code=item["code"],
                defaults={
                    "name": item["description"][:240],
                    "description": item["description"],
                    "unit": item["unit"] or "Stk.",
                    "purchase_price": item["purchase_price"] or Decimal("0"),
                    "sales_price": item["sales_price"] or item["purchase_price"] or Decimal("0"),
                    "supplier": supplier_name,
                    "external_codes": {"price_source_id": source.pk},
                    "active": True,
                },
            )
    return {"source": source.name, "created": created, "rows": len(items), "sha256": digest}


def import_reference_tree(root: str | Path, organization: Organization) -> dict:
    root = Path(root)
    results = []
    errors = []
    if not root.exists():
        return {"files": 0, "rows": 0, "results": [], "errors": [f"Verzeichnis nicht gefunden: {root}"]}
    for path in sorted(p for p in root.rglob("*") if p.is_file() and p.suffix.lower() in PRICE_EXTENSIONS):
        try:
            results.append(import_price_file(path, organization))
        except Exception as exc:
            errors.append({"file": str(path), "error": str(exc)})
    return {"files": len(results), "rows": sum(item["rows"] for item in results), "results": results, "errors": errors}
