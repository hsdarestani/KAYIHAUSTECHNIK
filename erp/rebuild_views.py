from __future__ import annotations

import base64
import calendar as month_calendar
import datetime as dt
import csv
import io
import json
import os
import re
import uuid
from datetime import timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path
from urllib.parse import urlencode

from django import forms
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.files.base import ContentFile
from django.db import transaction
from django.db.models import Count, Max, Q, Sum
from django.http import JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.urls import reverse
from django.views.decorators.http import require_http_methods, require_POST

from . import models as m
from .services.invoice_compliance_service import ComplianceError, finalize_invoice, get_compliance, is_finalized, save_customer_profile
from .services.business_pdf_identity import inject_business_pdf_identity
from .services.effective_pricing import catalog_with_effective_prices


# A+Bau intentionally keeps the existing ERP data model. This module replaces
# the legacy interaction model with a ToolTime-like operational flow while the
# old routes remain available as a fallback during migration.


def _org(request):
    profile = getattr(request.user, "profile", None)
    organization = getattr(profile, "organization", None)
    if organization is not None:
        return organization
    organization = m.Organization.objects.first()
    if organization is None:
        organization = m.Organization.objects.create(name="A+Bau")
    return organization


def _employee(request, org=None):
    org = org or _org(request)
    employee = getattr(request.user, "employee", None)
    if employee is not None and employee.organization_id == org.id:
        return employee
    email = getattr(request.user, "email", "") or ""
    if email:
        employee = m.Employee.objects.filter(organization=org, email__iexact=email).first()
    return employee


def _role(request):
    profile = getattr(request.user, "profile", None)
    return getattr(profile, "role", "office") or "office"


def _is_field_user(request):
    return _role(request) == "technician" or bool(getattr(getattr(request.user, "profile", None), "is_mobile_worker", False))


def _unique_number(model, org, prefix):
    year = timezone.localdate().year
    base = f"{prefix}-{year}-"
    latest = (
        model.objects.filter(organization=org, number__startswith=base)
        .order_by("-number")
        .values_list("number", flat=True)
        .first()
    )
    value = 1
    if latest:
        match = re.search(r"(\d+)$", latest)
        if match:
            value = int(match.group(1)) + 1
    while True:
        number = f"{base}{value:04d}"
        if not model.objects.filter(organization=org, number=number).exists():
            return number
        value += 1


def _money(value):
    try:
        return Decimal(str(value or "0").replace(",", "."))
    except (InvalidOperation, ValueError):
        return Decimal("0")


def _commercial_settings(document):
    if document is None:
        return None
    try:
        return document.commercial_settings
    except (m.CommercialDocumentSettings.DoesNotExist, AttributeError):
        return None


def _item_commercial_meta(item):
    try:
        return item.commercial_meta
    except (m.CommercialItemMeta.DoesNotExist, AttributeError):
        return None


def _discount_amount(net, settings, legacy_percent=Decimal("0")):
    if settings is None:
        value = net * _money(legacy_percent) / Decimal("100")
        return min(net, max(Decimal("0"), value))
    value = max(Decimal("0"), _money(settings.discount_value))
    if settings.discount_type == "fixed":
        return min(net, value)
    return min(net, net * value / Decimal("100"))


def _document_totals(document, *, include_payments=False):
    if document is None:
        base = {"net": Decimal("0"), "tax": Decimal("0"), "gross": Decimal("0"), "cost": Decimal("0"), "margin": Decimal("0"), "margin_percent": Decimal("0"), "discount": Decimal("0"), "alternative": Decimal("0"), "contingent": Decimal("0")}
        if include_payments:
            base.update({"paid": Decimal("0"), "open": Decimal("0")})
        return base
    settings = _commercial_settings(document)
    normal = alternative = contingent = cost = Decimal("0")
    for item in document.items.all():
        line = _money(item.quantity) * _money(item.unit_price)
        meta = _item_commercial_meta(item)
        model = getattr(meta, "service_model", "normal") if meta else "normal"
        if model == "alternative":
            alternative += line
            continue
        if model == "contingent":
            contingent += line
            continue
        normal += line
        purchase = _money(getattr(meta, "purchase_price", 0)) if meta else Decimal("0")
        if purchase <= 0 and getattr(item, "catalog_item_id", None):
            purchase = _money(getattr(item.catalog_item, "purchase_price", 0))
        cost += _money(item.quantity) * purchase
    legacy_discount = getattr(document, "discount_percent", Decimal("0"))
    discount = _discount_amount(normal, settings, legacy_discount)
    net = max(Decimal("0"), normal - discount)
    tax_rate = _money(getattr(settings, "tax_rate", 19) if settings else 19)
    tax = net * tax_rate / Decimal("100")
    gross = net + tax
    margin = net - cost
    margin_percent = (margin / net * Decimal("100")) if net else Decimal("0")
    result = {"net": net, "tax": tax, "gross": gross, "cost": cost, "margin": margin, "margin_percent": margin_percent, "discount": discount, "alternative": alternative, "contingent": contingent}
    if include_payments:
        paid = document.payments.aggregate(total=Sum("amount"))["total"] or Decimal("0")
        result.update({"paid": paid, "open": max(Decimal("0"), gross - paid)})
    return result


def _project_total(project):
    return sum((_invoice_total(invoice)["gross"] for invoice in project.invoices.exclude(status="cancelled").prefetch_related("items", "items__catalog_item", "items__commercial_meta", "payments")), Decimal("0"))


def _quote_total(quote):
    return _document_totals(quote, include_payments=False)


def _invoice_total(invoice):
    return _document_totals(invoice, include_payments=True)


def _project_financials(project):
    invoices = list(project.invoices.exclude(status="cancelled").prefetch_related("items", "items__catalog_item", "items__commercial_meta", "payments"))
    quotes = list(project.quotes.exclude(status="rejected").prefetch_related("items", "items__catalog_item", "items__commercial_meta"))
    invoice_totals = [_invoice_total(invoice) for invoice in invoices]
    quote_totals = [_quote_total(quote) for quote in quotes]
    revenue = sum((row["net"] for row in invoice_totals), Decimal("0"))
    costs = sum((row["cost"] for row in invoice_totals), Decimal("0"))
    margin = revenue - costs
    return {
        "revenue": revenue,
        "cost": costs,
        "margin": margin,
        "margin_percent": (margin / revenue * Decimal("100")) if revenue else Decimal("0"),
        "gross": sum((row["gross"] for row in invoice_totals), Decimal("0")),
        "paid": sum((row["paid"] for row in invoice_totals), Decimal("0")),
        "open": sum((row["open"] for row in invoice_totals), Decimal("0")),
        "quote_volume": sum((row["net"] for row in quote_totals), Decimal("0")),
    }

GERMAN_FIELD_LABELS = {
    "type": "Typ", "company": "Firma", "salutation": "Anrede", "first_name": "Vorname",
    "last_name": "Nachname", "email": "E-Mail", "phone": "Telefon", "mobile": "Mobil",
    "street": "Straße", "postal_code": "PLZ", "city": "Ort", "country": "Land",
    "vat_id": "USt-IdNr.", "notes": "Notizen", "title": "Titel", "customer": "Kunde",
    "object_location": "Objekt / Einsatzort", "description": "Beschreibung", "priority": "Priorität",
    "manager": "Projektleitung", "members": "Team", "starts_at": "Beginn", "ends_at": "Ende",
    "all_day": "Ganztägig", "location": "Einsatzort", "project": "Projekt", "attendees": "Mitarbeiter",
    "issue_date": "Ausstellungsdatum", "valid_until": "Gültig bis", "intro_text": "Einleitungstext",
    "outro_text": "Schlusstext", "discount_percent": "Rabatt (%)", "quote": "Angebot",
    "due_date": "Fällig am", "service_date": "Leistungsdatum", "status": "Status",
    "assigned_to": "Zugewiesen an", "due_at": "Fällig am", "supplier": "Lieferant",
    "amount_net": "Netto-Betrag", "tax_rate": "MwSt. (%)", "expense_date": "Belegdatum",
    "category": "Kategorie", "paid": "Bezahlt", "document": "Beleg / Dokument",
    "trade": "Gewerk", "hourly_cost": "Interner Stundensatz", "hourly_rate": "Verrechnungssatz",
    "active": "Aktiv", "color": "Farbe", "name": "Name",
}


class StyledModelForm(forms.ModelForm):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        for name, field in self.fields.items():
            if name in GERMAN_FIELD_LABELS:
                field.label = GERMAN_FIELD_LABELS[name]
            field.widget.attrs.setdefault("lang", "de")
            existing = field.widget.attrs.get("class", "")
            if isinstance(field.widget, forms.CheckboxInput):
                classes = [part for part in existing.split() if part not in {"next-control", "nx-control"}]
                classes.append("nx-checkbox-input")
                field.widget.attrs["class"] = " ".join(dict.fromkeys(classes))
            else:
                field.widget.attrs["class"] = f"{existing} next-control".strip()
            if isinstance(field.widget, forms.Textarea):
                field.widget.attrs.setdefault("rows", 4)


class CustomerForm(StyledModelForm):
    customer_number = forms.CharField(label="Kundennummer", required=False, max_length=30)

    class Meta:
        model = m.Customer
        fields = [
            "type", "company", "salutation", "first_name", "last_name", "email", "phone", "mobile",
            "street", "postal_code", "city", "country", "vat_id", "debtor_number", "routing_id", "supplier_id", "notes",
        ]
        widgets = {"notes": forms.Textarea(attrs={"rows": 3})}

    def __init__(self, *args, organization=None, modal=False, **kwargs):
        super().__init__(*args, **kwargs)
        self.organization = organization
        labels = {
            "type": "Kundentyp",
            "company": "Firmenname",
            "salutation": "Anrede",
            "first_name": "Vorname",
            "last_name": "Nachname",
            "email": "E-Mail",
            "phone": "Telefon",
            "mobile": "Mobil",
            "street": "Straße und Hausnummer",
            "postal_code": "PLZ",
            "city": "Ort",
            "country": "Land",
            "vat_id": "USt-IdNr.",
            "debtor_number": "Debitorennummer",
            "routing_id": "Routing-ID",
            "supplier_id": "Lieferanten-ID",
            "notes": "Beschreibung",
        }
        for name, label in labels.items():
            if name in self.fields:
                self.fields[name].label = label
        self.fields["customer_number"].widget.attrs.setdefault("placeholder", "Optional")
        if getattr(self.instance, "pk", None):
            self.fields["customer_number"].initial = self.instance.number
        if modal:
            self.fields["type"].choices = [("business", "Firmenkunde"), ("private", "Privatkunde")]
            self.fields["type"].widget = forms.RadioSelect(choices=self.fields["type"].choices)
        for name in ("debtor_number", "routing_id", "supplier_id", "vat_id"):
            self.fields[name].required = False
            self.fields[name].widget.attrs.setdefault("placeholder", "Optional")


class ObjectLocationForm(StyledModelForm):
    class Meta:
        model = m.ObjectLocation
        fields = ["name", "street", "postal_code", "city", "floor", "access_notes"]
        widgets = {"access_notes": forms.Textarea(attrs={"rows": 2})}


class ProjectForm(StyledModelForm):
    class Meta:
        model = m.Project
        fields = ["title", "customer", "object_location", "description", "priority", "manager", "members"]
        widgets = {"description": forms.Textarea(attrs={"rows": 3}), "members": forms.SelectMultiple(attrs={"size": 5})}

    def __init__(self, *args, organization=None, **kwargs):
        super().__init__(*args, **kwargs)
        self.fields["title"].label = "Projekttitel"
        self.fields["customer"].label = "Kunde auswählen"
        self.fields["object_location"].label = "Ausführungsort"
        self.fields["description"].label = "Projektbeschreibung"
        self.fields["priority"].label = "Priorität"
        self.fields["manager"].label = "Projektleitung"
        self.fields["members"].label = "Mitarbeiter"
        self.fields["customer"].widget.attrs["data-searchable"] = "true"
        self.fields["customer"].widget.attrs["data-search-placeholder"] = "Kunde suchen …"
        self.fields["object_location"].empty_label = "Kundenadresse verwenden (Standard)"
        self.fields["priority"].required = False
        self.fields["manager"].required = False
        self.fields["members"].required = False
        self.fields["description"].required = False
        self.fields["priority"].initial = self.initial.get("priority") or "normal"
        if organization:
            self.fields["customer"].queryset = m.Customer.objects.filter(organization=organization, active=True)
            self.fields["manager"].queryset = m.Employee.objects.filter(organization=organization, active=True)
            self.fields["members"].queryset = m.Employee.objects.filter(organization=organization, active=True)

            customer_id = None
            if self.is_bound:
                customer_id = self.data.get("customer")
            if not customer_id:
                customer_id = self.initial.get("customer")
            if not customer_id and getattr(self.instance, "customer_id", None):
                customer_id = self.instance.customer_id
            try:
                customer_id = int(customer_id) if customer_id else None
            except (TypeError, ValueError):
                customer_id = None
            locations = m.ObjectLocation.objects.filter(organization=organization)
            self.fields["object_location"].queryset = locations.filter(customer_id=customer_id) if customer_id else locations.none()

    def clean_priority(self):
        return self.cleaned_data.get("priority") or "normal"


class AppointmentForm(StyledModelForm):
    class Meta:
        model = m.CalendarEvent
        fields = ["title", "type", "starts_at", "ends_at", "all_day", "location", "notes", "project", "attendees"]
        widgets = {
            "starts_at": forms.DateTimeInput(attrs={"type": "datetime-local"}, format="%Y-%m-%dT%H:%M"),
            "ends_at": forms.DateTimeInput(attrs={"type": "datetime-local"}, format="%Y-%m-%dT%H:%M"),
            "notes": forms.Textarea(attrs={"rows": 3}),
            "attendees": forms.SelectMultiple(attrs={"size": 5}),
        }

    def __init__(self, *args, organization=None, **kwargs):
        super().__init__(*args, **kwargs)
        self.fields["starts_at"].input_formats = ["%Y-%m-%dT%H:%M"]
        self.fields["ends_at"].input_formats = ["%Y-%m-%dT%H:%M"]
        self.fields["project"].widget.attrs.update({"data-appointment-project": "1", "data-searchable": "true"})
        self.fields["attendees"].widget.attrs.update({"data-appointment-team": "1", "data-searchable": "true"})
        if organization:
            projects = (
                m.Project.objects.filter(organization=organization, archived=False)
                .select_related("customer", "object_location")
                .order_by("-updated_at")
            )
            self.fields["project"].queryset = projects
            self.fields["attendees"].queryset = m.Employee.objects.filter(
                organization=organization, active=True
            ).order_by("last_name", "first_name")
            # Non-model presentation querysets: no migration and no duplicate
            # persistence path. They only drive ToolTime-like customer/project UX.
            self.appointment_projects = projects
            self.appointment_customers = m.Customer.objects.filter(
                organization=organization, active=True
            ).order_by("company", "last_name", "first_name", "number")


class QuoteForm(StyledModelForm):
    class Meta:
        model = m.Quote
        fields = ["project", "issue_date", "valid_until", "intro_text", "discount_percent", "notes"]
        widgets = {
            "issue_date": forms.DateInput(attrs={"type": "date"}),
            "valid_until": forms.DateInput(attrs={"type": "date"}),
            "intro_text": forms.Textarea(attrs={"rows": 2}),
            "outro_text": forms.Textarea(attrs={"rows": 2}),
            "notes": forms.Textarea(attrs={"rows": 2}),
        }

    def __init__(self, *args, organization=None, **kwargs):
        super().__init__(*args, **kwargs)
        if organization:
            self.fields["project"].queryset = m.Project.objects.filter(organization=organization, archived=False)


class InvoiceForm(StyledModelForm):
    class Meta:
        model = m.Invoice
        fields = ["project", "quote", "issue_date", "service_date", "intro_text", "notes"]
        widgets = {
            "issue_date": forms.DateInput(attrs={"type": "date"}),
            "due_date": forms.DateInput(attrs={"type": "date"}),
            "service_date": forms.DateInput(attrs={"type": "date"}),
            "intro_text": forms.Textarea(attrs={"rows": 2}),
            "outro_text": forms.Textarea(attrs={"rows": 2}),
            "notes": forms.Textarea(attrs={"rows": 2}),
        }

    def __init__(self, *args, organization=None, **kwargs):
        super().__init__(*args, **kwargs)
        if organization:
            self.fields["project"].queryset = m.Project.objects.filter(organization=organization, archived=False)
            self.fields["quote"].queryset = m.Quote.objects.filter(organization=organization)


@login_required
def dashboard(request):
    if _is_field_user(request):
        return redirect("next-field")
    org = _org(request)
    today = timezone.localdate()
    start = timezone.make_aware(timezone.datetime.combine(today, timezone.datetime.min.time()))
    end = start + timedelta(days=1)
    appointments = (
        m.CalendarEvent.objects.filter(organization=org, starts_at__gte=start, starts_at__lt=end)
        .select_related("project", "project__customer")
        .prefetch_related("attendees")
        .order_by("starts_at")[:8]
    )
    open_quotes = m.Quote.objects.filter(organization=org, status__in=["draft", "review", "sent"]).count()
    overdue = m.Invoice.objects.filter(organization=org, status="overdue").count()
    active_projects = m.Project.objects.filter(organization=org, archived=False).exclude(status__in=["completed", "cancelled"]).count()
    recent_projects = m.Project.objects.filter(organization=org, archived=False).select_related("customer").order_by("-updated_at")[:6]
    return render(request, "rebuild/dashboard.html", {
        "appointments": appointments,
        "open_quotes": open_quotes,
        "overdue": overdue,
        "active_projects": active_projects,
        "recent_projects": recent_projects,
        "today": today,
    })


@login_required
def customer_list(request):
    org = _org(request)
    return render(request, "rebuild/customers.html", _customer_list_context(request, org))


def _customer_list_context(request, org, create_form=None, create_location_form=None, modal_open=False):
    keyword = (request.GET.get("keyword") or request.GET.get("q") or "").strip()
    sort_type = (request.GET.get("sortType") or "NAME").upper()
    if sort_type not in {"NAME", "LAST_CHANGE", "PROJECTS"}:
        sort_type = "NAME"
    sort_order = (request.GET.get("sortOrder") or "ASCENDING").upper()
    if sort_order not in {"ASCENDING", "DESCENDING"}:
        sort_order = "ASCENDING"
    try:
        offset = max(0, int(request.GET.get("offset") or 0))
    except (TypeError, ValueError):
        offset = 0

    customers_qs = m.Customer.objects.filter(organization=org, active=True).annotate(projects_count=Count("projects", distinct=True))
    if keyword:
        customers_qs = customers_qs.filter(
            Q(company__icontains=keyword)
            | Q(first_name__icontains=keyword)
            | Q(last_name__icontains=keyword)
            | Q(email__icontains=keyword)
            | Q(phone__icontains=keyword)
            | Q(mobile__icontains=keyword)
            | Q(number__icontains=keyword)
            | Q(debtor_number__icontains=keyword)
            | Q(routing_id__icontains=keyword)
            | Q(supplier_id__icontains=keyword)
            | Q(street__icontains=keyword)
            | Q(city__icontains=keyword)
        )

    descending = sort_order == "DESCENDING"
    prefix = "-" if descending else ""
    if sort_type == "LAST_CHANGE":
        ordering = [f"{prefix}updated_at", f"{prefix}pk"]
    elif sort_type == "PROJECTS":
        ordering = [f"{prefix}projects_count", "company", "last_name", "first_name", "pk"]
    else:
        ordering = [f"{prefix}company", f"{prefix}last_name", f"{prefix}first_name", f"{prefix}pk"]

    total_count = customers_qs.count()
    page_size = 50
    if total_count:
        max_offset = ((total_count - 1) // page_size) * page_size
        offset = min(offset, max_offset)
    else:
        offset = 0
    customers = list(customers_qs.order_by(*ordering)[offset : offset + page_size])

    if create_form is None:
        create_form = CustomerForm(
            organization=org,
            modal=True,
            initial={"type": "business", "customer_number": _unique_number(m.Customer, org, "K")},
        )
    if create_location_form is None:
        create_location_form = ObjectLocationForm(prefix="site")

    return {
        "customers": customers,
        "keyword": keyword,
        "query": keyword,
        "sort_type": sort_type,
        "sort_order": sort_order,
        "offset": offset,
        "page_size": page_size,
        "total_count": total_count,
        "prev_offset": max(0, offset - page_size),
        "next_offset": offset + page_size,
        "has_prev": offset > 0,
        "has_next": offset + page_size < total_count,
        "name_sort_next": "DESCENDING" if sort_type == "NAME" and sort_order == "ASCENDING" else "ASCENDING",
        "last_change_sort_next": "DESCENDING" if sort_type == "LAST_CHANGE" and sort_order == "ASCENDING" else "ASCENDING",
        "projects_sort_next": "DESCENDING" if sort_type == "PROJECTS" and sort_order == "ASCENDING" else "ASCENDING",
        "create_form": create_form,
        "create_location_form": create_location_form,
        "customer_modal_open": modal_open,
    }


@login_required
@require_http_methods(["GET", "POST"])
def customer_create(request):
    org = _org(request)
    modal = request.GET.get("modal") == "1" or request.POST.get("modal") == "1"
    next_target = (request.POST.get("next") if request.method == "POST" else request.GET.get("next")) or ""
    next_target = next_target if next_target in {"project"} else ""
    suggested_number = _unique_number(m.Customer, org, "K")
    form = CustomerForm(
        request.POST if request.method == "POST" else None,
        organization=org,
        modal=modal,
        initial={"type": "business", "customer_number": suggested_number},
    )
    location_payload = request.POST if request.method == "POST" and str(request.POST.get("site-street") or "").strip() else None
    location_form = ObjectLocationForm(location_payload, prefix="site")
    if request.method == "POST" and form.is_valid():
        requested_number = str(form.cleaned_data.get("customer_number") or "").strip() or suggested_number
        if m.Customer.objects.filter(organization=org, number=requested_number).exists():
            form.add_error("customer_number", "Diese Kundennummer ist bereits vergeben.")
        else:
            location_requested = bool(str(request.POST.get("site-street") or "").strip())
            location_valid = (not location_requested) or location_form.is_valid()
            if location_valid:
                with transaction.atomic():
                    customer = form.save(commit=False)
                    customer.organization = org
                    customer.number = requested_number
                    customer.save()
                    if location_requested:
                        location = location_form.save(commit=False)
                        location.organization = org
                        location.customer = customer
                        location.save()
                messages.success(request, "Kunde wurde angelegt.")
                if next_target == "project":
                    return redirect(f"/projects/new/?customer={customer.pk}")
                return redirect("next-customer-detail", pk=customer.pk)

    if modal:
        context = _customer_list_context(
            request,
            org,
            create_form=form,
            create_location_form=location_form,
            modal_open=True,
        )
        return render(request, "rebuild/customers.html", context)
    return render(request, "rebuild/customer_form.html", {
        "form": form,
        "location_form": location_form,
        "mode": "create",
        "next_target": next_target,
    })

def _customer_direct_document_project(org, customer):
    title = f"Direktdokumente · Kunde {customer.pk}"
    project = m.Project.objects.filter(organization=org, customer=customer, title=title).order_by("pk").first()
    if project is None:
        project = m.Project.objects.create(
            organization=org,
            customer=customer,
            number=_unique_number(m.Project, org, "P"),
            title=title,
            status="inquiry",
            archived=True,
        )
    return project


def _customer_bind_document_meta(document, kind, customer):
    if not hasattr(m, "ToolTimeDocumentMeta"):
        return
    lookup = {"organization": document.organization, kind: document}
    meta, _ = m.ToolTimeDocumentMeta.objects.get_or_create(**lookup)
    if hasattr(meta, "customer_id"):
        meta.customer = customer
        update_fields = ["customer"]
        if hasattr(meta, "updated_at"):
            update_fields.append("updated_at")
        meta.save(update_fields=update_fields)


@login_required
@require_POST
def customer_quote_create(request, pk):
    org = _org(request)
    if _is_field_user(request):
        messages.error(request, "Angebote können nur im Büro erstellt werden.")
        return redirect("next-customer-detail", pk=pk)
    customer = get_object_or_404(m.Customer, pk=pk, organization=org, active=True)
    project = _customer_direct_document_project(org, customer)
    quote = m.Quote.objects.create(
        organization=org,
        project=project,
        number="",
        status="draft",
        issue_date=timezone.localdate(),
        discount_percent=0,
        created_by=request.user,
    )
    _customer_bind_document_meta(quote, "quote", customer)
    return redirect("next-quote-edit", pk=quote.pk)


@login_required
@require_POST
def customer_invoice_create(request, pk):
    org = _org(request)
    if _is_field_user(request):
        messages.error(request, "Rechnungen können nur im Büro erstellt werden.")
        return redirect("next-customer-detail", pk=pk)
    customer = get_object_or_404(m.Customer, pk=pk, organization=org, active=True)
    project = _customer_direct_document_project(org, customer)
    today = timezone.localdate()
    invoice = m.Invoice.objects.create(
        organization=org,
        project=project,
        number="",
        status="draft",
        issue_date=today,
        due_date=today + timedelta(days=14),
        service_date=today,
        created_by=request.user,
    )
    _customer_bind_document_meta(invoice, "invoice", customer)
    return redirect("next-invoice-edit", pk=invoice.pk)


@login_required
@require_http_methods(["GET", "POST"])
def customer_detail(request, pk):
    """ToolTime-style customer cockpit while keeping A+Bau's existing data model."""
    org = _org(request)
    customer = get_object_or_404(m.Customer, pk=pk, organization=org)
    field_user = _is_field_user(request)
    active_tab = request.GET.get("tab", "overview")
    if active_tab not in {"overview", "tasks", "documents"}:
        active_tab = "overview"
    edit_mode = request.GET.get("edit") == "1"
    add_object_open = request.GET.get("add_object") == "1"

    if request.method == "POST" and request.POST.get("action") == "add_location":
        form = CustomerForm(instance=customer)
        location_form = ObjectLocationForm(request.POST, prefix="site")
        add_object_open = True
        if location_form.is_valid():
            location = location_form.save(commit=False)
            location.organization = org
            location.customer = customer
            location.save()
            messages.success(request, "Einsatzort wurde hinzugefügt.")
            return redirect(f"/customers/{customer.pk}/#weitere-kontakte")
    elif request.method == "POST":
        form = CustomerForm(request.POST, instance=customer)
        location_form = ObjectLocationForm(prefix="site")
        edit_mode = True
        if form.is_valid():
            form.save()
            messages.success(request, "Kundendaten gespeichert.")
            return redirect("next-customer-detail", pk=customer.pk)
    else:
        form = CustomerForm(instance=customer)
        location_form = ObjectLocationForm(prefix="site")

    direct_title = f"Direktdokumente · Kunde {customer.pk}"
    projects = (
        customer.projects.filter(organization=org)
        .exclude(title=direct_title)
        .select_related("object_location", "manager")
        .order_by("-updated_at")
    )
    locations = customer.object_locations.filter(organization=org).order_by("name", "city", "street")
    appointments = (
        m.CalendarEvent.objects.filter(organization=org)
        .filter(Q(customer=customer) | Q(project__customer=customer))
        .distinct()
        .select_related("project")
        .prefetch_related("attendees")
        .order_by("-starts_at")
    )
    quotes = (
        m.Quote.objects.filter(organization=org, project__customer=customer)
        .select_related("project", "created_by")
        .prefetch_related("items")
        .order_by("-issue_date", "-created_at")
    )
    invoices = (
        m.Invoice.objects.filter(organization=org, project__customer=customer)
        .select_related("project", "created_by")
        .prefetch_related("items", "payments")
        .order_by("-issue_date", "-created_at")
    )
    expenses = (
        m.Expense.objects.filter(organization=org).filter(Q(project__customer=customer) | Q(document__customer=customer)).distinct()
        .select_related("project", "document")
        .order_by("-expense_date", "-created_at")
    )
    tasks = (
        m.Task.objects.filter(project__customer=customer)
        .select_related("project")
        .order_by("status", "due_at", "-created_at")
    )
    documents = (
        m.Document.objects.filter(organization=org)
        .filter(Q(customer=customer) | Q(project__customer=customer))
        .select_related("project", "uploaded_by")
        .distinct()
        .order_by("-created_at")
    )

    def document_totals(document, with_payments=False):
        net = Decimal("0")
        tax = Decimal("0")
        for item in document.items.all():
            line_net = (item.quantity or Decimal("0")) * (item.unit_price or Decimal("0"))
            net += line_net
            tax += line_net * (item.tax_rate or Decimal("0")) / Decimal("100")
        if hasattr(document, "discount_percent") and document.discount_percent:
            factor = (Decimal("100") - document.discount_percent) / Decimal("100")
            net *= factor
            tax *= factor
        gross = net + tax
        paid = Decimal("0")
        if with_payments:
            paid = sum((payment.amount or Decimal("0") for payment in document.payments.all()), Decimal("0"))
        return {
            "net": net,
            "tax": tax,
            "gross": gross,
            "paid": paid,
            "open": max(Decimal("0"), gross - paid),
        }

    quote_rows = []
    for quote in quotes:
        try:
            totals = _quote_total(quote)
        except Exception:
            totals = document_totals(quote)
        quote.tt_totals = totals
        quote_rows.append(quote)

    invoice_rows = []
    revenue_net = Decimal("0")
    open_invoice_gross = Decimal("0")
    for invoice in invoices:
        try:
            totals = _invoice_total(invoice)
        except Exception:
            totals = document_totals(invoice, with_payments=True)
        invoice.tt_totals = totals
        invoice_rows.append(invoice)
        if invoice.status != "cancelled":
            revenue_net += totals["net"]
            open_invoice_gross += totals["open"]

    expenditure_net = sum((expense.amount_net or Decimal("0") for expense in expenses), Decimal("0"))

    # AB_V3_HISTORY_CUSTOMER_CONTEXT
    documented_event_ids = set(
        m.Document.objects.filter(
            organization=org,
            customer=customer,
            category="report",
            metadata__event_id__isnull=False,
        ).values_list("metadata__event_id", flat=True)
    )
    documented_event_ids.update(
        m.Document.objects.filter(
            organization=org,
            project__customer=customer,
            category="report",
            metadata__event_id__isnull=False,
        ).values_list("metadata__event_id", flat=True)
    )
    history = []
    for project in projects[:100]:
        history.append({"kind": "project", "at": project.updated_at, "object": project, "status": project.get_status_display()})
    for event in appointments[:100]:
        documented = event.pk in documented_event_ids
        history.append({
            "kind": "appointment",
            "at": event.starts_at,
            "object": event,
            "status": "Termin dokumentiert" if documented else "Termin geplant",
            "documented": documented,
        })
    for quote in quotes[:100]:
        history.append({"kind": "quote", "at": quote.created_at, "object": quote, "status": quote.get_status_display()})
    for invoice in invoices[:100]:
        history.append({"kind": "invoice", "at": invoice.created_at, "object": invoice, "status": invoice.get_status_display()})
    for expense in m.Expense.objects.filter(organization=org, project__customer=customer).order_by("-pk")[:100]:
        history.append({"kind": "expense", "at": getattr(expense, "created_at", timezone.now()), "object": expense, "status": expense.supplier or "Ausgabe"})
    history.sort(key=lambda row: row["at"], reverse=True)

    return render(request, "rebuild/customer_detail.html", {
        "customer": customer,
        "form": form,
        "projects": projects,
        "locations": locations,
        "location_form": location_form,
        "add_object_open": add_object_open,
        "edit_mode": edit_mode,
        "active_tab": active_tab,
        "appointments": appointments,
        "quotes": quote_rows,
        "invoices": invoice_rows,
        "expenses": expenses,
        "tasks": tasks,
        "documents": documents,
        "revenue_net": revenue_net,
        "expenditure_net": expenditure_net,
        "open_invoice_gross": open_invoice_gross,
        "history": history,
        "field_user": field_user,
        "direct_document_project_title": direct_title,
    })

@login_required
def customer_locations_api(request, pk):
    org = _org(request)
    customer = get_object_or_404(m.Customer, pk=pk, organization=org, active=True)
    locations = customer.object_locations.filter(organization=org).order_by("name", "city", "street")
    return JsonResponse({
        "customer": {
            "id": customer.pk,
            "name": customer.display_name,
            "address": ", ".join(part for part in [customer.street, f"{customer.postal_code} {customer.city}".strip()] if part),
        },
        "locations": [
            {
                "id": location.pk,
                "name": location.name or "Einsatzort",
                "address": ", ".join(part for part in [location.street, f"{location.postal_code} {location.city}".strip()] if part),
            }
            for location in locations
        ],
    })

@login_required
def supplier_list(request):
    org = _org(request)
    keyword = (request.GET.get("keyword") or request.GET.get("q") or "").strip()
    suppliers = m.CatalogItem.objects.filter(organization=org, active=True).exclude(supplier="")
    if keyword:
        suppliers = suppliers.filter(supplier__icontains=keyword)
    suppliers = suppliers.values("supplier").annotate(items_count=Count("id"), last_change=Max("updated_at")).order_by("supplier")
    return render(request, "rebuild/suppliers.html", {"suppliers": suppliers[:500], "keyword": keyword})

@login_required
def project_list(request):
    org = _org(request)
    query = request.GET.get("q", "").strip()
    status = request.GET.get("status", "").strip()
    projects = m.Project.objects.filter(organization=org, archived=False).select_related("customer", "manager", "object_location")
    if query:
        projects = projects.filter(Q(number__icontains=query) | Q(title__icontains=query) | Q(customer__company__icontains=query) | Q(customer__last_name__icontains=query))
    if status:
        projects = projects.filter(status=status)
    return render(request, "rebuild/projects.html", {"projects": projects.order_by("-updated_at")[:250], "query": query, "status": status, "statuses": m.Project._meta.get_field("status").choices})


@login_required
@require_http_methods(["GET", "POST"])
def project_create(request):
    # Technicians must use the price-free field intake; the office project form
    # must never be a bypass around approval/pricing controls.
    if _is_field_user(request):
        return redirect("field-quick-job")

    org = _org(request)
    source_id = (
        request.POST.get("_appointment")
        if request.method == "POST"
        else request.GET.get("_appointment") or request.GET.get("appointment") or request.GET.get("appointment_id")
    )
    source_appointment = None
    if source_id:
        try:
            source_id = int(source_id)
        except (TypeError, ValueError):
            source_id = None
        if source_id:
            source_appointment = m.CalendarEvent.objects.filter(
                organization=org, pk=source_id, project__isnull=True
            ).first()

    initial = {"priority": "normal"}
    if source_appointment is not None:
        initial["title"] = source_appointment.title
        initial["description"] = source_appointment.notes
    requested_customer = request.GET.get("customer")
    if requested_customer and m.Customer.objects.filter(
        organization=org, active=True, pk=requested_customer
    ).exists():
        initial["customer"] = requested_customer

    data = request.POST.copy() if request.method == "POST" else None
    if data is not None and not data.get("priority"):
        data["priority"] = "normal"
    form = ProjectForm(data, organization=org, initial=initial)

    if request.method == "POST" and form.is_valid():
        project = form.save(commit=False)
        project.organization = org
        project.number = _unique_number(m.Project, org, "P")
        project.status = "inquiry"
        project.save()
        form.save_m2m()
        if source_appointment is not None:
            source_appointment.project = project
            source_appointment.save()
            messages.success(
                request,
                "Projekt angelegt und Termin automatisch zugeordnet. "
                "Kundenfreigabe und Einsatzdokumentation sind jetzt verfügbar.",
            )
            return redirect("next-appointment-detail", pk=source_appointment.pk)
        messages.success(request, "Projekt wurde angelegt.")
        return redirect("next-project-detail", pk=project.pk)

    return render(request, "rebuild/project_form.html", {
        "form": form,
        "source_appointment": source_appointment,
    })


@login_required
def project_detail(request, pk):
    org = _org(request)
    project = get_object_or_404(m.Project.objects.select_related("customer", "object_location", "manager"), pk=pk, organization=org)
    appointments = project.events.prefetch_related("attendees").order_by("-starts_at")[:20]
    quotes = project.quotes.prefetch_related("items").order_by("-created_at")
    invoices = project.invoices.prefetch_related("items", "payments").order_by("-created_at")
    documents = project.documents.order_by("-created_at")[:30]
    tasks = project.tasks.order_by("status", "due_at")[:20]
    materials = project.materials.order_by("-created_at")[:20]
    invoice_gross = sum((_invoice_total(invoice)["gross"] for invoice in invoices), Decimal("0"))
    # KAYI_TOOLTIME_PROJECT_DETAIL_SAFE_CONTEXT
    quote_rows = [{"quote": quote, "total": _quote_total(quote)} for quote in quotes]
    invoice_rows = [{"invoice": invoice, "total": _invoice_total(invoice)} for invoice in invoices]
    turnover_net = sum((row["total"]["net"] for row in invoice_rows), Decimal("0"))
    open_amount = sum((row["total"]["open"] for row in invoice_rows), Decimal("0"))
    project_expenditure = Decimal("0")
    expense_model = getattr(m, "Expense", None)
    if expense_model is not None:
        try:
            fields = {field.name for field in expense_model._meta.get_fields()}
            if "project" in fields:
                expense_qs = expense_model.objects.filter(project=project)
                if "organization" in fields:
                    expense_qs = expense_qs.filter(organization=org)
                for expense in expense_qs[:500]:
                    raw = None
                    for attr in ("net_amount", "amount_net", "net", "amount"):
                        if hasattr(expense, attr):
                            raw = getattr(expense, attr)
                            if raw is not None:
                                break
                    project_expenditure += Decimal(str(raw or 0))
        except Exception:
            project_expenditure = Decimal("0")
    receipts = [
        document for document in documents
        if str(getattr(document, "category", "")).lower() in {"receipt", "expense", "beleg"}
        or str((getattr(document, "metadata", {}) or {}).get("kind", "")).lower() in {"receipt", "expense", "beleg"}
    ]
    field_user = _is_field_user(request)

    # AB_V3_HISTORY_PROJECT_CONTEXT
    documented_event_ids = set(
        m.Document.objects.filter(
            organization=org,
            project=project,
            category="report",
            metadata__event_id__isnull=False,
        ).values_list("metadata__event_id", flat=True)
    )
    history = []
    for event in appointments:
        documented = event.pk in documented_event_ids
        history.append({
            "kind": "appointment",
            "at": event.starts_at,
            "object": event,
            "status": "Termin dokumentiert" if documented else "Termin geplant",
            "documented": documented,
        })
    for quote in quotes:
        history.append({"kind": "quote", "at": quote.created_at, "object": quote, "status": quote.get_status_display()})
    for invoice in invoices:
        history.append({"kind": "invoice", "at": invoice.created_at, "object": invoice, "status": invoice.get_status_display()})
    for document in documents:
        history.append({"kind": "document", "at": document.created_at, "object": document, "status": document.get_category_display()})
    history.sort(key=lambda row: row["at"], reverse=True)

    latest_invoice = invoices[0] if invoices else None
    latest_quote = quotes[0] if quotes else None
    if project.status == "completed":
        tooltime_status = "Projekt abgeschlossen"
    elif project.status == "cancelled":
        tooltime_status = "Projekt abgebrochen"
    elif latest_invoice is not None:
        if latest_invoice.status == "paid":
            tooltime_status = "Rechnung bezahlt"
        elif latest_invoice.status == "overdue":
            tooltime_status = "Rechnung überfällig"
        elif latest_invoice.status in {"dunned", "reminded"}:
            tooltime_status = "Rechnung angemahnt"
        else:
            tooltime_status = "Rechnung angelegt"
    elif latest_quote is not None:
        tooltime_status = "Angebot angenommen" if latest_quote.status == "accepted" else "Angebot erstellt"
    elif documented_event_ids:
        tooltime_status = "Termin dokumentiert"
    elif appointments:
        tooltime_status = "Termin geplant"
    else:
        tooltime_status = "Neues Projekt"
    has_planned_appointments = any(event.pk not in documented_event_ids for event in appointments)

    return render(request, "rebuild/project_detail.html", {
        "history": history,
        "tooltime_status": tooltime_status,
        "has_planned_appointments": has_planned_appointments,
        "project": project, "appointments": appointments, "quotes": quotes, "invoices": invoices,
        "documents": documents, "tasks": tasks, "materials": materials, "invoice_gross": invoice_gross,
        "quote_rows": quote_rows,
        "invoice_rows": invoice_rows,
        "turnover_net": turnover_net,
        "project_expenditure": project_expenditure,
        "open_amount": open_amount,
        "receipts": receipts,
        "field_user": field_user,
    })


@login_required
@require_POST
def project_lifecycle(request, pk):
    org = _org(request)
    if _is_field_user(request):
        messages.error(request, "Projektstatus kann nur im Büro geändert werden.")
        return redirect("next-project-detail", pk=pk)
    project = get_object_or_404(m.Project, pk=pk, organization=org)
    action = (request.POST.get("action") or "").strip()

    if action in {"complete", "cancel"}:
        event_ids = list(project.events.values_list("pk", flat=True))
        documented = set(
            m.Document.objects.filter(
                organization=org,
                project=project,
                category="report",
                metadata__event_id__in=event_ids,
            ).values_list("metadata__event_id", flat=True)
        )
        if any(event_id not in documented for event_id in event_ids):
            messages.error(
                request,
                "Das Projekt kann erst abgeschlossen oder abgebrochen werden, wenn keine geplanten Termine mehr enthalten sind.",
            )
            return redirect("next-project-detail", pk=project.pk)
        project.status = "completed" if action == "complete" else "cancelled"
        project.archived = True
        update_fields = ["status", "archived", "updated_at"]
        if action == "complete" and hasattr(project, "progress"):
            project.progress = 100
            update_fields.append("progress")
        project.save(update_fields=update_fields)
        messages.success(
            request,
            "Projekt wurde abgeschlossen." if action == "complete" else "Projekt wurde abgebrochen.",
        )
    elif action == "reactivate":
        project.status = "inquiry"
        project.archived = False
        project.save(update_fields=["status", "archived", "updated_at"])
        messages.success(request, "Projekt wurde reaktiviert.")
    else:
        messages.error(request, "Unbekannte Projektaktion.")
    return redirect("next-project-detail", pk=project.pk)


@login_required
def appointment_list(request):
    org = _org(request)
    allowed_views = {"day", "week", "month", "map", "list"}
    calendar_view = (request.GET.get("view") or "week").strip().lower()
    if calendar_view not in allowed_views:
        calendar_view = "week"

    raw_date = (request.GET.get("date") or request.GET.get("start") or "").strip()
    try:
        anchor_date = timezone.datetime.strptime(raw_date, "%Y-%m-%d").date() if raw_date else timezone.localdate()
    except ValueError:
        anchor_date = timezone.localdate()

    employee_filter = (request.GET.get("employee") or "").strip()
    project_filter = (request.GET.get("project") or "").strip()
    query_filter = (request.GET.get("q") or "").strip()
    customer_filter = (request.GET.get("customer") or "").strip()
    events = (
        m.CalendarEvent.objects.filter(organization=org)
        .select_related("project", "project__customer", "customer")
        .prefetch_related("attendees")
    )
    if employee_filter.isdigit():
        events = events.filter(attendees__pk=int(employee_filter)).distinct()
    else:
        employee_filter = ""
    if project_filter.isdigit():
        events = events.filter(project_id=int(project_filter))
    else:
        project_filter = ""
    if customer_filter.isdigit():
        events = events.filter(
            Q(project__customer_id=int(customer_filter)) | Q(customer_id=int(customer_filter))
        ).distinct()
    else:
        customer_filter = ""
    if query_filter:
        events = events.filter(
            Q(title__icontains=query_filter)
            | Q(location__icontains=query_filter)
            | Q(project__number__icontains=query_filter)
            | Q(project__title__icontains=query_filter)
            | Q(project__customer__company__icontains=query_filter)
            | Q(project__customer__first_name__icontains=query_filter)
            | Q(project__customer__last_name__icontains=query_filter)
            | Q(customer__company__icontains=query_filter)
            | Q(customer__first_name__icontains=query_filter)
            | Q(customer__last_name__icontains=query_filter)
        ).distinct()

    def month_shift(value, delta):
        index = value.year * 12 + (value.month - 1) + delta
        year, month_index = divmod(index, 12)
        return value.replace(year=year, month=month_index + 1, day=1)

    month_first = anchor_date.replace(day=1)
    next_month_first = month_shift(month_first, 1)
    week_start = anchor_date - timedelta(days=anchor_date.weekday())

    if calendar_view == "day":
        range_start = anchor_date
        range_end = anchor_date + timedelta(days=1)
        prev_date = anchor_date - timedelta(days=1)
        next_date = anchor_date + timedelta(days=1)
    elif calendar_view == "month":
        range_start = month_first - timedelta(days=month_first.weekday())
        month_last = next_month_first - timedelta(days=1)
        range_end = month_last + timedelta(days=(6 - month_last.weekday()) + 1)
        prev_date = month_shift(month_first, -1)
        next_date = next_month_first
    elif calendar_view in {"list", "map"}:
        range_start = anchor_date
        range_end = anchor_date + timedelta(days=90)
        prev_date = anchor_date - timedelta(days=30)
        next_date = anchor_date + timedelta(days=30)
    else:
        range_start = week_start
        range_end = week_start + timedelta(days=7)
        prev_date = week_start - timedelta(days=7)
        next_date = week_start + timedelta(days=7)

    tz = timezone.get_current_timezone()
    aware_start = timezone.make_aware(timezone.datetime.combine(range_start, timezone.datetime.min.time()), tz)
    aware_end = timezone.make_aware(timezone.datetime.combine(range_end, timezone.datetime.min.time()), tz)
    event_list = list(events.filter(starts_at__lt=aware_end, ends_at__gte=aware_start).order_by("starts_at"))

    grouped = {}
    for event in event_list:
        local_start = timezone.localtime(event.starts_at)
        event.ui_date = local_start.date()
        event_customer = event.project.customer if event.project_id and event.project.customer_id else event.customer
        event.ui_customer = event_customer.display_name if event_customer is not None else "Intern"
        event.ui_project = f"{event.project.number} · {event.project.title}" if event.project_id else ("Ohne Projekt" if event.customer_id else "Interner Termin")
        event.ui_map_address = (event.location or "").strip() or _appointment_customer_address(event_customer)
        attendee_names = []
        for attendee in event.attendees.all():
            name = f"{attendee.first_name} {attendee.last_name}".strip()
            attendee_names.append(name or getattr(attendee, "email", "") or str(attendee.pk))
        event.ui_attendees = ", ".join(attendee_names) if attendee_names else "Nicht zugewiesen"
        grouped.setdefault(event.ui_date, []).append(event)

    def calendar_url(target_date, target_view=None):
        params = {"view": target_view or calendar_view, "date": target_date.isoformat()}
        if employee_filter:
            params["employee"] = employee_filter
        if project_filter:
            params["project"] = project_filter
        if customer_filter:
            params["customer"] = customer_filter
        if query_filter:
            params["q"] = query_filter
        return "?" + urlencode(params)

    today = timezone.localdate()
    days = []
    month_weeks = []
    list_days = []
    day_events = grouped.get(anchor_date, [])
    map_events = [event for event in event_list if event.ui_map_address]

    if calendar_view == "week":
        for offset in range(7):
            day_date = week_start + timedelta(days=offset)
            days.append({
                "date": day_date,
                "events": grouped.get(day_date, []),
                "is_today": day_date == today,
                "day_url": calendar_url(day_date, "day"),
            })
    elif calendar_view == "month":
        cursor = range_start
        while cursor < range_end:
            week = []
            for _ in range(7):
                day_events_for_date = grouped.get(cursor, [])
                week.append({
                    "date": cursor,
                    "events": day_events_for_date[:4],
                    "extra_count": max(0, len(day_events_for_date) - 4),
                    "in_month": cursor.month == month_first.month,
                    "is_today": cursor == today,
                    "day_url": calendar_url(cursor, "day"),
                })
                cursor += timedelta(days=1)
            month_weeks.append(week)
    elif calendar_view == "list":
        cursor = range_start
        while cursor < range_end:
            date_events = grouped.get(cursor, [])
            if date_events:
                list_days.append({"date": cursor, "events": date_events, "is_today": cursor == today})
            cursor += timedelta(days=1)

    employees = m.Employee.objects.filter(organization=org, active=True).order_by("last_name", "first_name")
    projects = m.Project.objects.filter(organization=org, archived=False).select_related("customer").order_by("-updated_at")[:300]
    customers = m.Customer.objects.filter(organization=org, active=True).order_by("company", "last_name", "first_name")[:300]
    view_links = [
        {"key": "day", "label": "Tag", "url": calendar_url(anchor_date, "day")},
        {"key": "week", "label": "Woche", "url": calendar_url(anchor_date, "week")},
        {"key": "month", "label": "Monat", "url": calendar_url(anchor_date, "month")},
        {"key": "map", "label": "Karte", "url": calendar_url(anchor_date, "map")},
        {"key": "list", "label": "Liste", "url": calendar_url(anchor_date, "list")},
    ]
    reset_query = "?" + urlencode({"view": calendar_view, "date": anchor_date.isoformat()})

    return render(request, "rebuild/appointments.html", {
        "calendar_view": calendar_view,
        "anchor_date": anchor_date,
        "range_start": range_start,
        "range_end": range_end - timedelta(days=1),
        "month_first": month_first,
        "week_start": week_start,
        "days": days,
        "month_weeks": month_weeks,
        "day_events": day_events,
        "map_events": map_events,
        "list_days": list_days,
        "view_links": view_links,
        "prev_url": calendar_url(prev_date),
        "next_url": calendar_url(next_date),
        "today_url": calendar_url(today),
        "employees": employees,
        "projects": projects,
        "customers": customers,
        "employee_filter": employee_filter,
        "project_filter": project_filter,
        "customer_filter": customer_filter,
        "query_filter": query_filter,
        "reset_query": reset_query,
        "can_dispatch": not _is_field_user(request),
    })


@login_required
@require_POST
def appointment_move(request, pk):
    org = _org(request)
    if _is_field_user(request):
        return JsonResponse({"ok": False, "error": "Terminverschiebungen sind nur für Büro/Leitung verfügbar."}, status=403)
    event = get_object_or_404(m.CalendarEvent, pk=pk, organization=org)
    try:
        payload = json.loads(request.body.decode("utf-8") or "{}")
    except (json.JSONDecodeError, UnicodeDecodeError):
        return JsonResponse({"ok": False, "error": "Ungültige Anfrage."}, status=400)
    raw_date = str(payload.get("date") or "").strip()
    try:
        target_date = timezone.datetime.strptime(raw_date, "%Y-%m-%d").date()
    except ValueError:
        return JsonResponse({"ok": False, "error": "Bitte ein gültiges Datum wählen."}, status=400)

    local_start = timezone.localtime(event.starts_at)
    duration = event.ends_at - event.starts_at
    local_time = local_start.time().replace(tzinfo=None)
    target_start = timezone.make_aware(
        timezone.datetime.combine(target_date, local_time),
        timezone.get_current_timezone(),
    )
    event.starts_at = target_start
    event.ends_at = target_start + duration
    event.save(update_fields=["starts_at", "ends_at", "updated_at"])
    return JsonResponse({
        "ok": True,
        "id": event.pk,
        "date": target_date.isoformat(),
        "starts_at": timezone.localtime(event.starts_at).isoformat(),
        "ends_at": timezone.localtime(event.ends_at).isoformat(),
    })


def _appointment_recurrence_request(request):
    if request.method != "POST":
        return "none", 4, 1, "day", None
    rule = (request.POST.get("repeat_rule") or "none").strip().lower()
    allowed = {"none", "daily", "weekdays", "weekly", "biweekly", "monthly", "half_yearly", "yearly", "custom"}
    if rule not in allowed:
        rule = "none"
    try:
        count = int(request.POST.get("repeat_count") or 4)
    except (TypeError, ValueError):
        count = 4
    count = max(2, min(count, 52)) if rule != "none" else 1
    try:
        interval = int(request.POST.get("repeat_interval") or 1)
    except (TypeError, ValueError):
        interval = 1
    interval = max(1, min(interval, 52))
    unit = (request.POST.get("repeat_unit") or "day").strip().lower()
    if unit not in {"day", "weekday", "week", "month", "year"}:
        unit = "day"
    until = None
    end_mode = (request.POST.get("repeat_end_mode") or "count").strip().lower()
    until_raw = (request.POST.get("repeat_until") or "").strip()
    if rule == "custom" and end_mode == "date" and until_raw:
        try:
            until = dt.date.fromisoformat(until_raw)
        except ValueError:
            until = None
    return rule, count, interval, unit, until


def _appointment_shift_month(value, months):
    local_value = timezone.localtime(value)
    month_index = local_value.year * 12 + (local_value.month - 1) + months
    year, zero_based_month = divmod(month_index, 12)
    month = zero_based_month + 1
    day = min(local_value.day, month_calendar.monthrange(year, month)[1])
    return local_value.replace(year=year, month=month, day=day)


def _appointment_shift_weekdays(value, index):
    if index <= 0:
        return value
    shifted = value
    remaining = index
    while remaining:
        shifted = shifted + timedelta(days=1)
        if timezone.localtime(shifted).weekday() < 5:
            remaining -= 1
    return shifted


def _appointment_recurrence_shift(value, rule, index, interval=1, unit="day"):
    if rule == "custom":
        step = interval * index
        if unit == "day":
            return value + timedelta(days=step)
        if unit == "weekday":
            return _appointment_shift_weekdays(value, step)
        if unit == "week":
            return value + timedelta(days=7 * step)
        if unit == "month":
            return _appointment_shift_month(value, step)
        if unit == "year":
            return _appointment_shift_month(value, 12 * step)
    if rule == "daily":
        return value + timedelta(days=index)
    if rule == "weekdays":
        return _appointment_shift_weekdays(value, index)
    if rule == "weekly":
        return value + timedelta(days=7 * index)
    if rule == "biweekly":
        return value + timedelta(days=14 * index)
    if rule == "monthly":
        return _appointment_shift_month(value, index)
    if rule == "half_yearly":
        return _appointment_shift_month(value, 6 * index)
    if rule == "yearly":
        return _appointment_shift_month(value, 12 * index)
    return value


def _appointment_recurrence_indices(starts_at, rule, count, interval=1, unit="day", until=None):
    if rule == "none":
        return []
    if rule == "custom" and until is not None:
        indices = []
        for index in range(1, 731):
            candidate = _appointment_recurrence_shift(starts_at, rule, index, interval, unit)
            if timezone.localtime(candidate).date() > until:
                break
            indices.append(index)
        return indices
    return list(range(1, count))


def _appointment_customer_address(customer):
    if customer is None:
        return ""
    city_line = " ".join(part for part in [customer.postal_code, customer.city] if part).strip()
    return ", ".join(part for part in [customer.street, city_line] if part)


def _appointment_event_customer(event):
    if event.project_id and getattr(event.project, "customer_id", None):
        return event.project.customer
    return getattr(event, "customer", None)


def _appointment_source_quote(request, org):
    raw = (request.POST.get("source_quote") if request.method == "POST" else request.GET.get("quote")) or ""
    raw = str(raw).strip()
    if not raw.isdigit():
        return None
    return m.Quote.objects.filter(
        organization=org, pk=int(raw), status="accepted"
    ).select_related("project__customer").prefetch_related("items").first()


def _appointment_seed_groups(*, event=None, quote=None, request=None):
    if request is not None and request.method == "POST" and request.POST.get("service_editor_present") == "1":
        titles = request.POST.getlist("service_group_title")
        group_indexes = request.POST.getlist("service_group_index")
        descriptions = request.POST.getlist("service_description")
        kinds = request.POST.getlist("service_kind")
        quantities = request.POST.getlist("service_quantity")
        units = request.POST.getlist("service_unit")
        catalog_ids = request.POST.getlist("service_catalog_id")
        purchase_prices = request.POST.getlist("service_purchase_price")
        unit_prices = request.POST.getlist("service_unit_price")
        taxes = request.POST.getlist("service_tax_rate")
        mixed_payloads = request.POST.getlist("service_mixed_json")
        source_ids = request.POST.getlist("service_source_quote_item_id")
        result = []
        for group_index, title in enumerate(titles):
            items = []
            for index, description in enumerate(descriptions):
                if index >= len(group_indexes) or str(group_indexes[index]) != str(group_index):
                    continue
                description = (description or "").strip()
                if not description:
                    continue
                items.append({
                    "kind": kinds[index] if index < len(kinds) else "other",
                    "quantity": quantities[index] if index < len(quantities) else "1",
                    "unit": units[index] if index < len(units) else "Stk.",
                    "description": description,
                    "catalog_item_id": catalog_ids[index] if index < len(catalog_ids) else "",
                    "purchase_price": purchase_prices[index] if index < len(purchase_prices) else "0",
                    "unit_price": unit_prices[index] if index < len(unit_prices) else "0",
                    "tax_rate": taxes[index] if index < len(taxes) else "19",
                    "mixed_json": mixed_payloads[index] if index < len(mixed_payloads) else "[]",
                    "source_quote_item_id": source_ids[index] if index < len(source_ids) else "",
                })
            result.append({"title": (title or "").strip(), "items": items})
        return result
    if event is not None:
        result = []
        for group in event.service_groups.prefetch_related("items__catalog_item").all().order_by("position", "id"):
            items = []
            for item in group.items.all().order_by("position", "id"):
                items.append({
                    "kind": item.kind, "quantity": str(item.quantity), "unit": item.unit,
                    "description": item.description, "catalog_item_id": item.catalog_item_id or "",
                    "purchase_price": str(item.purchase_price), "unit_price": str(item.unit_price),
                    "tax_rate": str(item.tax_rate),
                    "mixed_json": json.dumps(item.mixed_payload or [], ensure_ascii=False),
                    "source_quote_item_id": item.source_quote_item_id or "",
                })
            result.append({"title": group.title, "items": items})
        return result
    if quote is not None:
        grouped, order = {}, []
        items = quote.items.select_related("catalog_item").prefetch_related("tooltime_mixed_subitems").order_by("position", "pk")
        for item in items:
            try:
                meta = item.commercial_meta
            except Exception:
                meta = None
            title = (getattr(meta, "group_title", "") or "").strip()
            key = title or "__default__"
            if key not in grouped:
                grouped[key] = []
                order.append(key)
            mixed = [{
                "item_type": sub.item_type, "description": sub.description,
                "quantity": str(sub.quantity), "unit": sub.unit,
                "purchase_price": str(sub.purchase_price), "sales_price": str(sub.sales_price),
            } for sub in item.tooltime_mixed_subitems.all().order_by("sort_order", "id")]
            grouped[key].append({
                "kind": getattr(meta, "position_type", "other") if meta else "other",
                "quantity": str(item.quantity), "unit": item.unit, "description": item.description,
                "catalog_item_id": item.catalog_item_id or "",
                "purchase_price": str(getattr(meta, "purchase_price", 0) or 0),
                "unit_price": str(item.unit_price), "tax_rate": str(item.tax_rate),
                "mixed_json": json.dumps(mixed, ensure_ascii=False), "source_quote_item_id": item.pk,
            })
        return [{"title": "" if key == "__default__" else key, "items": grouped[key]} for key in order]
    return []


def _appointment_save_services(event, request, source_quote=None):
    if request.POST.get("service_editor_present") != "1" and source_quote is None:
        return
    groups = _appointment_seed_groups(quote=source_quote, request=request)
    event.service_items.all().delete()
    event.service_groups.all().delete()
    for group_position, row in enumerate(groups, start=1):
        group = m.AppointmentServiceGroup.objects.create(
            organization=event.organization, event=event,
            title=(row.get("title") or "")[:220], position=group_position,
        )
        for item_position, row_item in enumerate(row.get("items") or [], start=1):
            catalog = None
            catalog_id = str(row_item.get("catalog_item_id") or "").strip()
            if catalog_id.isdigit():
                catalog = m.CatalogItem.objects.filter(
                    organization=event.organization, active=True, pk=int(catalog_id)
                ).first()
            source_item = None
            source_id = str(row_item.get("source_quote_item_id") or "").strip()
            if source_id.isdigit():
                source_item = m.QuoteItem.objects.filter(
                    quote__organization=event.organization, pk=int(source_id)
                ).first()
            kind = str(row_item.get("kind") or "other")
            if kind not in {"labour", "material", "mixed", "other"}:
                kind = "other"
            try:
                mixed = json.loads(row_item.get("mixed_json") or "[]")
            except (TypeError, ValueError, json.JSONDecodeError):
                mixed = []
            m.AppointmentServiceItem.objects.create(
                organization=event.organization, event=event, group=group,
                catalog_item=catalog, source_quote_item=source_item, position=item_position, kind=kind,
                code=(catalog.code if catalog else "")[:80],
                description=str(row_item.get("description") or "").strip(),
                quantity=_money(row_item.get("quantity") or 1),
                unit=str(row_item.get("unit") or (catalog.unit if catalog else "Stk."))[:30],
                purchase_price=max(Decimal("0"), _money(row_item.get("purchase_price") or (catalog.purchase_price if catalog else 0))),
                unit_price=max(Decimal("0"), _money(row_item.get("unit_price") or (catalog.sales_price if catalog else 0))),
                tax_rate=max(Decimal("0"), _money(row_item.get("tax_rate") or (catalog.tax_rate if catalog else 19))),
                mixed_payload=mixed if isinstance(mixed, list) else [],
            )


def _appointment_clone_services(source, target):
    target.service_items.all().delete()
    target.service_groups.all().delete()
    for group in source.service_groups.prefetch_related("items").all().order_by("position", "id"):
        target_group = m.AppointmentServiceGroup.objects.create(
            organization=target.organization, event=target, title=group.title, position=group.position
        )
        for item in group.items.all().order_by("position", "id"):
            m.AppointmentServiceItem.objects.create(
                organization=target.organization, event=target, group=target_group,
                catalog_item=item.catalog_item, source_quote_item=item.source_quote_item,
                position=item.position, kind=item.kind, code=item.code, description=item.description,
                quantity=item.quantity, unit=item.unit, purchase_price=item.purchase_price,
                unit_price=item.unit_price, tax_rate=item.tax_rate, mixed_payload=item.mixed_payload,
            )


def _appointment_service_snapshot(event):
    rows = []
    for group in event.service_groups.prefetch_related("items").all().order_by("position", "id"):
        for item in group.items.all().order_by("position", "id"):
            rows.append({
                "group": group.title, "kind": item.kind, "description": item.description,
                "quantity": str(item.quantity), "unit": item.unit, "catalog_item_id": item.catalog_item_id,
            })
    return rows


def _appointment_direct_project(event):
    if event.project_id:
        return event.project
    customer = _appointment_event_customer(event)
    if customer is None:
        return None
    title = f"Direktdokumente · Kunde {customer.pk}"
    project = m.Project.objects.filter(
        organization=event.organization, customer=customer, title=title
    ).order_by("pk").first()
    if project is None:
        project = m.Project.objects.create(
            organization=event.organization, customer=customer,
            number=_unique_number(m.Project, event.organization, "P"), title=title,
            status="inquiry", archived=True,
        )
    return project


def _appointment_bind_customer_meta(document, kind, customer):
    if customer is None or not hasattr(m, "ToolTimeDocumentMeta"):
        return
    lookup = {"organization": document.organization, kind: document}
    meta, _ = m.ToolTimeDocumentMeta.objects.get_or_create(**lookup)
    if hasattr(meta, "customer_id"):
        meta.customer = customer
        meta.save(update_fields=["customer", "updated_at"])


def _appointment_copy_services_to_document(event, document, kind):
    for item in event.service_items.select_related("catalog_item", "group").all().order_by("group__position", "position", "id"):
        kwargs = {
            "position": item.position, "description": item.description,
            "quantity": item.quantity, "unit": item.unit, "unit_price": item.unit_price,
            "tax_rate": item.tax_rate, "catalog_item": item.catalog_item,
        }
        kwargs[kind] = document
        target = (m.QuoteItem if kind == "quote" else m.InvoiceItem).objects.create(**kwargs)
        if hasattr(m, "CommercialItemMeta"):
            meta_kwargs = {
                "organization": event.organization,
                "position_type": item.kind,
                "purchase_price": item.purchase_price,
                "markup_percent": (((item.unit_price / item.purchase_price) - Decimal("1")) * Decimal("100")) if item.purchase_price > 0 else Decimal("0"),
                "group_title": item.group.title if item.group_id else "",
                ("quote_item" if kind == "quote" else "invoice_item"): target,
            }
            m.CommercialItemMeta.objects.create(**meta_kwargs)
        if item.mixed_payload and hasattr(m, "ToolTimeMixedSubitem"):
            for sub_index, sub in enumerate(item.mixed_payload):
                if not isinstance(sub, dict) or not str(sub.get("description") or "").strip():
                    continue
                sub_kwargs = {
                    "organization": event.organization,
                    "item_type": str(sub.get("item_type") or "other")[:16],
                    "description": str(sub.get("description") or "")[:300],
                    "quantity": _money(sub.get("quantity") or 1),
                    "unit": str(sub.get("unit") or "Stk.")[:30],
                    "purchase_price": max(Decimal("0"), _money(sub.get("purchase_price") or 0)),
                    "sales_price": max(Decimal("0"), _money(sub.get("sales_price") or 0)),
                    "sort_order": sub_index,
                    ("quote_item" if kind == "quote" else "invoice_item"): target,
                }
                m.ToolTimeMixedSubitem.objects.create(**sub_kwargs)


def _appointment_apply_field_services(event, request):
    ids = request.POST.getlist("document_service_id")
    kinds = request.POST.getlist("document_service_kind")
    quantities = request.POST.getlist("document_service_quantity")
    units = request.POST.getlist("document_service_unit")
    descriptions = request.POST.getlist("document_service_description")
    catalog_ids = request.POST.getlist("document_service_catalog_id")
    onsite_group = None
    for index, description in enumerate(descriptions):
        description = (description or "").strip()
        if not description:
            continue
        raw_id = ids[index] if index < len(ids) else ""
        item = event.service_items.filter(pk=int(raw_id)).first() if str(raw_id).isdigit() else None
        catalog = None
        raw_catalog = catalog_ids[index] if index < len(catalog_ids) else ""
        if str(raw_catalog).isdigit():
            catalog = m.CatalogItem.objects.filter(
                organization=event.organization, active=True, pk=int(raw_catalog)
            ).first()
        kind = kinds[index] if index < len(kinds) else "other"
        if catalog is not None:
            kind = "labour" if catalog.kind == "service" else ("material" if catalog.kind == "material" else "other")
        if kind not in {"labour", "material", "mixed", "other"}:
            kind = "other"
        if item is None:
            if onsite_group is None:
                onsite_group, _ = m.AppointmentServiceGroup.objects.get_or_create(
                    organization=event.organization, event=event, title="Vor Ort ergänzt",
                    defaults={"position": event.service_groups.count() + 1},
                )
            item = m.AppointmentServiceItem(
                organization=event.organization, event=event, group=onsite_group,
                position=onsite_group.items.count() + 1,
                purchase_price=catalog.purchase_price if catalog else 0,
                unit_price=catalog.sales_price if catalog else 0,
                tax_rate=catalog.tax_rate if catalog else 19,
            )
        elif catalog is not None and item.catalog_item_id != catalog.pk:
            item.purchase_price, item.unit_price, item.tax_rate = catalog.purchase_price, catalog.sales_price, catalog.tax_rate
        item.catalog_item = catalog or item.catalog_item
        item.kind = kind
        item.description = description
        item.quantity = _money(quantities[index] if index < len(quantities) else 1)
        item.unit = (units[index] if index < len(units) else (catalog.unit if catalog else "Stk.")) or "Stk."
        item.save()


@login_required
@require_http_methods(["GET", "POST"])
def appointment_create(request):
    org = _org(request)
    source_quote = _appointment_source_quote(request, org)
    now = timezone.localtime().replace(second=0, microsecond=0)
    initial = {"starts_at": now + timedelta(hours=1), "ends_at": now + timedelta(hours=2)}
    if source_quote is not None and source_quote.project_id and not request.GET.get("project"):
        initial["project"] = str(source_quote.project_id)
    requested_project_id = (
        request.POST.get("project") if request.method == "POST" else request.GET.get("project")
    )
    requested_customer_id = (
        request.POST.get("customer_filter") if request.method == "POST" else request.GET.get("customer")
    )

    # Resolve all preselection values through organization-scoped querysets.
    # Customer can be stored directly when no project is selected. A selected
    # project remains authoritative for its customer to prevent mismatched links.
    selected_project = None
    if requested_project_id:
        selected_project = m.Project.objects.filter(
            organization=org, archived=False, pk=requested_project_id
        ).select_related("customer").first()
        if selected_project is not None:
            initial["project"] = selected_project.pk

    selected_customer = None
    if requested_customer_id:
        selected_customer = m.Customer.objects.filter(
            organization=org, active=True, pk=requested_customer_id
        ).first()
    if selected_project is not None and selected_project.customer_id:
        selected_customer = selected_project.customer

    selected_customer_id = str(selected_customer.pk) if selected_customer is not None else ""
    selected_project_id = str(selected_project.pk) if selected_project is not None else ""
    repeat_rule, repeat_count, repeat_interval, repeat_unit, repeat_until = _appointment_recurrence_request(request)
    form = AppointmentForm(
        request.POST if request.method == "POST" else None,
        organization=org,
        initial=initial,
    )
    if request.method == "POST" and form.is_valid():
        event = form.save(commit=False)
        event.organization = org
        event.created_by = request.user
        event.source_quote = source_quote
        event.work_report = (request.POST.get("work_report") or "").strip()
        event.customer = selected_customer
        if not event.location and event.customer_id:
            event.location = _appointment_customer_address(event.customer)
        occurrence_indices = _appointment_recurrence_indices(
            event.starts_at, repeat_rule, repeat_count, repeat_interval, repeat_unit, repeat_until
        )
        series_id = uuid.uuid4() if repeat_rule != "none" and occurrence_indices else None
        event.recurrence_series = series_id
        event.recurrence_rule = repeat_rule
        event.recurrence_index = 0
        event.recurrence_interval = repeat_interval
        event.recurrence_unit = repeat_unit
        event.recurrence_until = repeat_until
        event.save()
        form.save_m2m()
        _appointment_save_services(event, request, source_quote)

        if series_id is not None:
            attendees = list(event.attendees.all())
            for occurrence_index in occurrence_indices:
                occurrence = m.CalendarEvent.objects.create(
                    organization=org,
                    customer=event.customer,
                    project=event.project,
                    title=event.title,
                    type=event.type,
                    starts_at=_appointment_recurrence_shift(event.starts_at, repeat_rule, occurrence_index, repeat_interval, repeat_unit),
                    ends_at=_appointment_recurrence_shift(event.ends_at, repeat_rule, occurrence_index, repeat_interval, repeat_unit),
                    all_day=event.all_day,
                    location=event.location,
                    notes=event.notes,
                    created_by=event.created_by,
                    source_quote=event.source_quote,
                    work_report=event.work_report,
                    recurrence_series=series_id,
                    recurrence_rule=repeat_rule,
                    recurrence_index=occurrence_index,
                    recurrence_interval=repeat_interval,
                    recurrence_unit=repeat_unit,
                    recurrence_until=repeat_until,
                )
                if attendees:
                    occurrence.attendees.set(attendees)
                _appointment_clone_services(event, occurrence)
            messages.success(request, f"Terminserie mit {1 + len(occurrence_indices)} Terminen wurde geplant.")
        else:
            messages.success(request, "Termin wurde geplant.")
        return redirect("next-appointment-detail", pk=event.pk)
    return render(request, "rebuild/appointment_form.html", {
        "form": form,
        "selected_customer_id": selected_customer_id,
        "selected_project_id": selected_project_id,
        "selected_repeat_rule": repeat_rule,
        "repeat_count": repeat_count,
        "repeat_interval": repeat_interval,
        "repeat_unit": repeat_unit,
        "repeat_until": repeat_until.isoformat() if repeat_until else "",
        "appointment_service_groups": _appointment_seed_groups(quote=source_quote, request=request),
        "appointment_catalog": m.CatalogItem.objects.filter(organization=org, active=True).order_by("name")[:500],
        "source_quote": source_quote,
    })


@login_required
@require_http_methods(["GET", "POST"])
def appointment_edit(request, pk):
    org = _org(request)
    event = get_object_or_404(
        m.CalendarEvent.objects.filter(organization=org)
        .select_related("project", "project__customer", "project__object_location", "customer")
        .prefetch_related("attendees"),
        pk=pk,
    )
    if _is_field_user(request):
        messages.warning(request, "Die Terminplanung kann nur im Büro bearbeitet werden.")
        return redirect("next-appointment-detail", pk=event.pk)

    requested_project_id = (
        request.POST.get("project") if request.method == "POST" else event.project_id
    )
    requested_customer_id = (
        request.POST.get("customer_filter") if request.method == "POST" else event.customer_id
    )

    selected_project = None
    if requested_project_id:
        selected_project = (
            m.Project.objects.filter(
                organization=org,
                archived=False,
                pk=requested_project_id,
            )
            .select_related("customer", "object_location")
            .first()
        )

    selected_customer = None
    if requested_customer_id:
        selected_customer = m.Customer.objects.filter(
            organization=org,
            active=True,
            pk=requested_customer_id,
        ).first()
    if selected_project is not None and selected_project.customer_id:
        selected_customer = selected_project.customer

    form = AppointmentForm(
        request.POST if request.method == "POST" else None,
        instance=event,
        organization=org,
    )
    if request.method == "POST" and form.is_valid():
        # ModelForm validation mutates the bound instance before this block.
        # Reload the persisted occurrence so the series time delta is measured
        # against the real pre-edit schedule rather than the already-cleaned form value.
        persisted_event = m.CalendarEvent.objects.get(organization=org, pk=event.pk)
        original_start = persisted_event.starts_at
        original_series = persisted_event.recurrence_series
        original_index = persisted_event.recurrence_index
        series_scope = (request.POST.get("series_scope") or "single").strip().lower()
        if series_scope not in {"single", "following", "all"}:
            series_scope = "single"

        updated = form.save(commit=False)
        updated.organization = org
        # Editing must never rewrite the original author merely because another
        # office user adjusts timing/team later.
        updated.created_by = event.created_by or request.user
        updated.customer = selected_customer
        if not updated.location and updated.customer_id:
            updated.location = _appointment_customer_address(updated.customer)
        updated.save()
        form.save_m2m()
        updated.work_report = (request.POST.get("work_report") or "").strip()
        updated.save(update_fields=["work_report", "updated_at"])
        _appointment_save_services(updated, request, None)

        affected = 1
        if original_series and series_scope in {"following", "all"}:
            targets = m.CalendarEvent.objects.filter(
                organization=org,
                recurrence_series=original_series,
            ).exclude(pk=updated.pk)
            if series_scope == "following":
                targets = targets.filter(recurrence_index__gte=original_index)

            start_delta = updated.starts_at - original_start
            duration = updated.ends_at - updated.starts_at
            attendees = list(updated.attendees.all())
            for occurrence in targets.order_by("recurrence_index"):
                occurrence.starts_at = occurrence.starts_at + start_delta
                occurrence.ends_at = occurrence.starts_at + duration
                occurrence.title = updated.title
                occurrence.type = updated.type
                occurrence.all_day = updated.all_day
                occurrence.location = updated.location
                occurrence.notes = updated.notes
                occurrence.work_report = updated.work_report
                occurrence.project = updated.project
                occurrence.customer = updated.customer
                occurrence.save()
                occurrence.attendees.set(attendees)
                _appointment_clone_services(updated, occurrence)
                affected += 1

        if affected > 1:
            label = "diesem und allen folgenden Terminen" if series_scope == "following" else "allen Terminen der Serie"
            messages.success(request, f"Änderungen wurden auf {label} übernommen ({affected} Termine).")
        else:
            messages.success(request, "Termin wurde aktualisiert.")
        return redirect("next-appointment-detail", pk=updated.pk)

    return render(request, "rebuild/appointment_form.html", {
        "form": form,
        "mode": "edit",
        "event": event,
        "selected_customer_id": str(selected_customer.pk) if selected_customer is not None else "",
        "selected_project_id": str(selected_project.pk) if selected_project is not None else "",
        "selected_repeat_rule": event.recurrence_rule,
        "repeat_count": 1,
        "repeat_interval": event.recurrence_interval,
        "repeat_unit": event.recurrence_unit,
        "repeat_until": event.recurrence_until.isoformat() if event.recurrence_until else "",
        "appointment_service_groups": _appointment_seed_groups(event=event, request=request),
        "appointment_catalog": m.CatalogItem.objects.filter(organization=org, active=True).order_by("name")[:500],
        "source_quote": event.source_quote,
    })


@login_required
@require_POST
def appointment_delete(request, pk):
    org = _org(request)
    event = get_object_or_404(m.CalendarEvent.objects.filter(organization=org), pk=pk)
    if _is_field_user(request):
        messages.warning(request, "Die Terminplanung kann nur im Büro gelöscht werden.")
        return redirect("next-appointment-detail", pk=event.pk)

    scope = (request.POST.get("scope") or "single").strip().lower()
    if scope == "following" and event.recurrence_series:
        deleted, _ = m.CalendarEvent.objects.filter(
            organization=org,
            recurrence_series=event.recurrence_series,
            recurrence_index__gte=event.recurrence_index,
        ).delete()
        messages.success(request, f"{deleted} Termin(e) der Serie wurden gelöscht.")
    else:
        event.delete()
        messages.success(request, "Termin wurde gelöscht.")
    return redirect("next-appointments")


@login_required
def appointment_detail(request, pk):
    org = _org(request)
    event = get_object_or_404(
        m.CalendarEvent.objects.select_related("project", "project__customer", "project__object_location", "customer", "source_quote"),
        pk=pk, organization=org,
    )
    docs = m.Document.objects.filter(organization=org, metadata__event_id=event.pk).order_by("-created_at")
    documented = docs.filter(category="report").exists()
    employee = _employee(request, org)
    running = None
    if employee and event.project:
        running = m.TimeEntry.objects.filter(
            organization=org, employee=employee, project=event.project, ended_at__isnull=True
        ).order_by("-started_at").first()
    return render(request, "rebuild/appointment_detail.html", {
        "event": event, "documents": docs, "documented": documented,
        "running": running, "employee": employee,
        "service_groups": event.service_groups.prefetch_related("items__catalog_item").all().order_by("position", "id"),
        "appointment_catalog": m.CatalogItem.objects.filter(organization=org, active=True).order_by("name")[:500],
        "event_customer": _appointment_event_customer(event),
    })


@login_required
def field_home(request):
    org = _org(request)
    employee = _employee(request, org)
    events = m.CalendarEvent.objects.filter(organization=org).select_related("project", "project__customer").prefetch_related("attendees")
    if employee:
        events = events.filter(attendees=employee)
    now = timezone.now()
    planned = events.filter(ends_at__gte=now).order_by("starts_at")[:30]
    overdue = events.filter(ends_at__lt=now).order_by("-starts_at")[:30]
    documented_ids = set(
        m.Document.objects.filter(organization=org, category="report", metadata__kind="field_completion", metadata__event_id__isnull=False)
        .exclude(metadata__status="changes_requested")
        .values_list("metadata__event_id", flat=True)
    )
    overdue = [event for event in overdue if event.pk not in documented_ids]
    documented = [event for event in events.order_by("-starts_at")[:100] if event.pk in documented_ids][:30]
    return render(request, "rebuild/field_home.html", {"planned": planned, "overdue": overdue, "documented": documented, "employee": employee})


@login_required
@require_POST
def time_toggle(request, event_pk):
    org = _org(request)
    event = get_object_or_404(m.CalendarEvent, pk=event_pk, organization=org)
    employee = _employee(request, org)
    if employee is None or event.project_id is None:
        return JsonResponse({"ok": False, "error": "Mitarbeiter oder Projekt fehlt."}, status=400)
    running = m.TimeEntry.objects.filter(organization=org, employee=employee, project=event.project, ended_at__isnull=True).order_by("-started_at").first()
    if running:
        running.ended_at = timezone.now()
        running.save(update_fields=["ended_at", "updated_at"])
        return JsonResponse({"ok": True, "state": "stopped"})
    entry = m.TimeEntry.objects.create(
        organization=org, employee=employee, project=event.project, started_at=timezone.now(),
        description=f"Termin #{event.pk}: {event.title}",
    )
    return JsonResponse({"ok": True, "state": "running", "id": entry.pk})


@login_required
@require_POST
def appointment_document(request, pk):
    org = _org(request)
    event = get_object_or_404(
        m.CalendarEvent.objects.select_related("project", "project__customer", "customer"), pk=pk, organization=org
    )
    customer = _appointment_event_customer(event)
    if customer is None:
        return JsonResponse({"ok": False, "error": "Dem Termin muss ein Kunde zugeordnet sein."}, status=400)
    _appointment_apply_field_services(event, request)
    report_text = (request.POST.get("report_text") or "").strip()
    services = (request.POST.get("services") or "").strip()
    material = (request.POST.get("material") or "").strip()
    customer_name = (request.POST.get("customer_name") or "").strip()
    body = report_text or event.work_report or "Vor-Ort-Dokumentation"
    payload = {
        "event_id": event.pk, "event_title": event.title,
        "services": services, "material": material, "customer_name": customer_name,
        "service_items": _appointment_service_snapshot(event), "source": "kayi-next-field",
    }
    report = m.Document(
        organization=org, customer=customer, project=event.project,
        title=f"Arbeitsbericht · {event.title} · {timezone.localdate():%d.%m.%Y}",
        category="report", mime_type="text/plain", size=len(body.encode("utf-8")),
        metadata=payload, uploaded_by=request.user,
    )
    report.file.save(f"arbeitsbericht-{event.pk}-{timezone.now():%Y%m%d%H%M%S}.txt", ContentFile(body.encode("utf-8")), save=False)
    report.save()
    for upload in request.FILES.getlist("photos"):
        photo = m.Document(
            organization=org, customer=customer, project=event.project,
            title=upload.name, category="photo", mime_type=getattr(upload, "content_type", "") or "",
            size=getattr(upload, "size", 0) or 0,
            metadata={"event_id": event.pk, "source": "kayi-next-field"}, uploaded_by=request.user,
        )
        photo.file.save(upload.name, upload, save=False)
        photo.save()
    signature_data = request.POST.get("signature_data") or ""
    if signature_data.startswith("data:image/png;base64,"):
        try:
            raw = base64.b64decode(signature_data.split(",", 1)[1])
            signature = m.Document(
                organization=org, customer=customer, project=event.project,
                title=f"Kundenunterschrift · {customer_name or customer.display_name}", category="other",
                mime_type="image/png", size=len(raw),
                metadata={"event_id": event.pk, "kind": "customer_signature"}, uploaded_by=request.user,
            )
            signature.file.save(f"signature-{event.pk}.png", ContentFile(raw), save=False)
            signature.save()
        except Exception:
            pass
    if event.project and event.project.status in {"inquiry", "planning", "quoted", "confirmed"}:
        event.project.status = "in_progress"
        event.project.actual_start = event.project.actual_start or timezone.localdate()
        event.project.save(update_fields=["status", "actual_start", "updated_at"])
    return JsonResponse({"ok": True, "redirect": f"/appointments/{event.pk}/"})


@login_required
@require_POST
def ai_structure_report(request, pk):
    org = _org(request)
    get_object_or_404(m.CalendarEvent, pk=pk, organization=org)
    from erp.store_views import has_ai_consent
    if not has_ai_consent(request.user):
        return JsonResponse({"ok": False, "error": "Vor der KI-Verarbeitung ist deine ausdrückliche Einwilligung in den Einstellungen erforderlich.", "consent_required": True, "settings_url": "/settings/next/"}, status=428)
    raw = (request.POST.get("text") or "").strip()
    if not raw:
        return JsonResponse({"ok": False, "error": "Kein Text vorhanden."}, status=400)
    fallback = {"report": raw, "services": "", "material": ""}
    key = os.environ.get("OPENAI_API_KEY", "").strip()
    if not key:
        return JsonResponse({"ok": True, **fallback, "ai": False})
    try:
        from openai import OpenAI
        client = OpenAI(api_key=key)
        response = client.responses.create(
            model=os.environ.get("OPENAI_MODEL", "gpt-4.1-mini"),
            input=[
                {"role": "system", "content": "Du strukturierst Baustellen-Diktate für einen deutschen Haustechnikbetrieb. Antworte ausschließlich als JSON mit report, services, material. Erfinde nichts."},
                {"role": "user", "content": raw},
            ],
        )
        text = response.output_text.strip()
        match = re.search(r"\{.*\}", text, re.S)
        data = json.loads(match.group(0) if match else text)
        return JsonResponse({"ok": True, "report": data.get("report", raw), "services": data.get("services", ""), "material": data.get("material", ""), "ai": True})
    except Exception:
        return JsonResponse({"ok": True, **fallback, "ai": False})


@login_required
@require_POST
def appointment_from_quote(request, pk):
    org = _org(request)
    if _is_field_user(request):
        messages.error(request, "Termine aus Angeboten können nur im Büro erstellt werden.")
        return redirect("next-quotes")
    quote = get_object_or_404(m.Quote.objects.select_related("project__customer"), organization=org, pk=pk)
    if quote.status != "accepted":
        messages.error(request, "Ein Termin kann erst aus einem angenommenen Angebot erstellt werden.")
        return redirect("next-quote-edit", pk=quote.pk)
    return redirect(f"{reverse('next-appointment-create')}?quote={quote.pk}")


@login_required
@require_POST
def appointment_to_quote(request, pk):
    org = _org(request)
    if _is_field_user(request):
        messages.error(request, "Angebote können nur im Büro erstellt werden.")
        return redirect("next-appointment-detail", pk=pk)
    event = get_object_or_404(
        m.CalendarEvent.objects.select_related("project__customer", "customer"), organization=org, pk=pk
    )
    customer = _appointment_event_customer(event)
    if customer is None:
        messages.error(request, "Für ein Angebot muss dem Termin ein Kunde zugeordnet sein.")
        return redirect("next-appointment-detail", pk=event.pk)
    quote = m.Quote.objects.create(
        organization=org, project=_appointment_direct_project(event), source_event=event,
        number="", status="draft", issue_date=timezone.localdate(), discount_percent=0, created_by=request.user,
    )
    _appointment_bind_customer_meta(quote, "quote", customer)
    _appointment_copy_services_to_document(event, quote, "quote")
    messages.success(request, "Angebotsentwurf wurde aus dem Termin erstellt. Alle Leistungen wurden übernommen.")
    return redirect("next-quote-edit", pk=quote.pk)


@login_required
@require_POST
def appointment_to_invoice(request, pk):
    org = _org(request)
    if _is_field_user(request):
        messages.error(request, "Rechnungen können nur im Büro erstellt werden.")
        return redirect("next-appointment-detail", pk=pk)
    event = get_object_or_404(
        m.CalendarEvent.objects.select_related("project__customer", "customer"), organization=org, pk=pk
    )
    customer = _appointment_event_customer(event)
    if customer is None:
        messages.error(request, "Für eine Rechnung muss dem Termin ein Kunde zugeordnet sein.")
        return redirect("next-appointment-detail", pk=event.pk)
    documented = m.Document.objects.filter(
        organization=org, category="report", metadata__event_id=event.pk
    ).exists()
    if not documented:
        messages.error(request, "Eine Rechnung kann erst aus einem dokumentierten Termin erstellt werden.")
        return redirect("next-appointment-detail", pk=event.pk)
    today = timezone.localdate()
    invoice = m.Invoice.objects.create(
        organization=org, project=_appointment_direct_project(event), source_event=event,
        number="", status="draft", issue_date=today, due_date=today + timedelta(days=14),
        service_date=today, created_by=request.user,
    )
    _appointment_bind_customer_meta(invoice, "invoice", customer)
    _appointment_copy_services_to_document(event, invoice, "invoice")
    messages.success(request, "Rechnungsentwurf wurde aus dem dokumentierten Termin erstellt. Alle Leistungen und gespeicherten Preise wurden übernommen.")
    return redirect("next-invoice-edit", pk=invoice.pk)


@login_required
def quote_list(request):
    org = _org(request)
    quotes = m.Quote.objects.filter(organization=org).select_related("project", "project__customer").prefetch_related("items").order_by("-created_at")[:250]
    rows = [{"quote": quote, "total": _quote_total(quote)} for quote in quotes]
    return render(request, "rebuild/quotes.html", {"rows": rows})


def _tax_from_code(code):
    return {"19": Decimal("19"), "7": Decimal("7"), "0_19": Decimal("0"), "0_13b": Decimal("0"), "0_4": Decimal("0"), "0": Decimal("0")}.get(code, Decimal("19"))


def _save_document_settings(document, request, kind):
    lookup = {"quote": document} if kind == "quote" else {"invoice": document}
    settings, _ = m.CommercialDocumentSettings.objects.get_or_create(organization=document.organization, **lookup)
    settings.tax_code = (request.POST.get("document_tax_code") or "19")[:20]
    settings.tax_rate = _tax_from_code(settings.tax_code)
    settings.discount_type = "fixed" if request.POST.get("discount_type") == "fixed" else "percent"
    settings.discount_value = max(Decimal("0"), _money(request.POST.get("discount_value")))
    settings.payment_due_days = max(0, min(365, int(_money(request.POST.get("payment_due_days") or 14))))
    settings.early_payment_discount_percent = max(Decimal("0"), min(Decimal("100"), _money(request.POST.get("early_discount_percent"))))
    settings.early_payment_discount_days = max(0, min(365, int(_money(request.POST.get("early_discount_days") or 0))))
    settings.closing_text = (request.POST.get("closing_text") or "").strip()
    settings.save()
    return settings


def _posted(values, index, default=""):
    return values[index] if index < len(values) else default


def _save_commercial_items(document, request, kind, settings):
    parent_field = "quote" if kind == "quote" else "invoice"
    model = m.QuoteItem if kind == "quote" else m.InvoiceItem
    descriptions = request.POST.getlist("item_description")
    detail_texts = request.POST.getlist("item_detail")
    quantities = request.POST.getlist("item_quantity")
    units = request.POST.getlist("item_unit")
    purchases = request.POST.getlist("item_purchase_price")
    markups = request.POST.getlist("item_markup_percent")
    item_types = request.POST.getlist("item_type")
    service_models = request.POST.getlist("item_service_model")
    catalog_ids = request.POST.getlist("item_catalog_id")
    groups = request.POST.getlist("item_group")
    sales_prices = request.POST.getlist("item_sales_price")
    mixed_payloads = request.POST.getlist("item_subitems_json")
    mixed_visibility = request.POST.getlist("item_show_subitems")
    document.items.all().delete()
    position = 1
    for index, raw_description in enumerate(descriptions):
        description = (raw_description or "").strip()
        if not description:
            continue
        quantity = _money(_posted(quantities, index, "1"))
        purchase_raw = (_posted(purchases, index, "") or "").strip()
        markup = max(Decimal("-100"), _money(_posted(markups, index, "0")))
        catalog = None
        catalog_id = (_posted(catalog_ids, index, "") or "").strip()
        if catalog_id.isdigit():
            catalog = m.CatalogItem.objects.filter(organization=document.organization, active=True, pk=int(catalog_id)).first()
        if purchase_raw:
            purchase = max(Decimal("0"), _money(purchase_raw))
        elif catalog is not None:
            purchase = max(Decimal("0"), _money(catalog.purchase_price))
            if purchase <= 0:
                purchase = max(Decimal("0"), _money(catalog.sales_price))
        else:
            purchase = Decimal("0")
        unit_price = (purchase * (Decimal("1") + markup / Decimal("100"))).quantize(Decimal("0.01"))
        manual_sales = (_posted(sales_prices, index, "") or "").strip()
        if manual_sales:
            unit_price = max(Decimal("0"), _money(manual_sales)).quantize(Decimal("0.01"))
            if purchase > 0:
                markup = ((unit_price / purchase) - Decimal("1")) * Decimal("100")
        if unit_price <= 0 and catalog is not None and _money(catalog.sales_price) > 0:
            unit_price = _money(catalog.sales_price).quantize(Decimal("0.01"))
            if purchase <= 0:
                purchase = unit_price
                markup = Decimal("0")
        kwargs = {
            parent_field: document,
            "position": position,
            "description": description,
            "quantity": quantity,
            "unit": (_posted(units, index, "Stk.") or "Stk.")[:30],
            "unit_price": unit_price,
            "tax_rate": settings.tax_rate,
            "catalog_item": catalog,
        }
        item = model.objects.create(**kwargs)
        detail = (_posted(detail_texts, index, "") or "").strip()
        position_type = (_posted(item_types, index, "material") or "material")
        if position_type not in {"material", "labour", "mixed", "other"}:
            position_type = "material"
        service_model = (_posted(service_models, index, "normal") or "normal")
        if service_model not in {"normal", "alternative", "contingent"}:
            service_model = "normal"
        meta_kwargs = {
            "organization": document.organization,
            "position_type": position_type,
            "purchase_price": purchase,
            "markup_percent": markup,
            "service_model": service_model,
            "detail_text": detail,
            "group_title": (_posted(groups, index, "") or "")[:220],
            "show_subitems_in_pdf": (_posted(mixed_visibility, index, "1") or "1") in {"1", "true", "on", "yes"},
        }
        meta_kwargs["quote_item" if kind == "quote" else "invoice_item"] = item
        item_meta = m.CommercialItemMeta.objects.create(**meta_kwargs)
        if position_type == "mixed":
            raw_payload = (_posted(mixed_payloads, index, "") or "").strip()
            try:
                payload = json.loads(raw_payload) if raw_payload else []
            except (TypeError, ValueError, json.JSONDecodeError):
                payload = []
            total_sales = Decimal("0")
            total_purchase = Decimal("0")
            for sub_index, row in enumerate(payload if isinstance(payload, list) else []):
                if not isinstance(row, dict):
                    continue
                sub_description = str(row.get("description") or "").strip()
                if not sub_description:
                    continue
                sub_type = str(row.get("item_type") or "material")
                if sub_type not in {"material", "labour", "other"}:
                    sub_type = "material"
                sub_quantity = _money(row.get("quantity") or 1)
                sub_purchase = max(Decimal("0"), _money(row.get("purchase_price") or 0))
                sub_sales = max(Decimal("0"), _money(row.get("sales_price") or 0))
                sub_kwargs = {
                    "organization": document.organization,
                    "item_type": sub_type,
                    "description": sub_description[:300],
                    "quantity": sub_quantity,
                    "unit": str(row.get("unit") or "Stk.")[:30],
                    "purchase_price": sub_purchase,
                    "sales_price": sub_sales,
                    "sort_order": sub_index,
                }
                sub_kwargs["quote_item" if kind == "quote" else "invoice_item"] = item
                m.ToolTimeMixedSubitem.objects.create(**sub_kwargs)
                total_sales += sub_quantity * sub_sales
                total_purchase += sub_quantity * sub_purchase
            if payload:
                item.unit_price = max(Decimal("0"), total_sales).quantize(Decimal("0.01"))
                item.save(update_fields=["unit_price"])
                item_meta.purchase_price = max(Decimal("0"), total_purchase).quantize(Decimal("0.01"))
                item_meta.markup_percent = (((item.unit_price / item_meta.purchase_price) - Decimal("1")) * Decimal("100")) if item_meta.purchase_price > 0 else Decimal("0")
                item_meta.save(update_fields=["purchase_price", "markup_percent", "updated_at"])
        if catalog is not None and detail and not (catalog.description or "").strip():
            catalog.description = detail
            catalog.save(update_fields=["description", "updated_at"])
        position += 1


def _items_for_editor(document):
    if document is None:
        return []
    rows = list(document.items.select_related("catalog_item").order_by("position"))
    for item in rows:
        meta = _item_commercial_meta(item)
        purchase = _money(getattr(meta, "purchase_price", 0)) if meta else Decimal("0")
        if purchase <= 0 and item.catalog_item_id:
            purchase = _money(item.catalog_item.purchase_price)
        if purchase <= 0:
            purchase = _money(item.unit_price)
        item.ui_purchase_price = purchase
        item.ui_markup_percent = _money(getattr(meta, "markup_percent", 0)) if meta else Decimal("0")
        item.ui_type = getattr(meta, "position_type", "material") if meta else ("material" if item.catalog_item_id else "other")
        item.ui_service_model = getattr(meta, "service_model", "normal") if meta else "normal"
        item.ui_detail = getattr(meta, "detail_text", "") if meta else (item.catalog_item.description if item.catalog_item_id else "")
        item.ui_group = getattr(meta, "group_title", "") if meta else ""
        item.ui_catalog_id = item.catalog_item_id or ""
        item.ui_show_subitems = getattr(meta, "show_subitems_in_pdf", True) if meta else True
        item.ui_mixed_subitems = list(item.tooltime_mixed_subitems.all().order_by("sort_order", "id"))
        item.ui_subitems_json = json.dumps([
            {"item_type": row.item_type, "description": row.description, "quantity": str(row.quantity), "unit": row.unit, "purchase_price": str(row.purchase_price), "sales_price": str(row.sales_price)}
            for row in item.ui_mixed_subitems
        ], ensure_ascii=False)
    return rows


def _fast_catalog_preview(org, limit=18):
    """Render only a small, already-priced local shortlist on first paint.

    Full B&O/VA04 discovery is intentionally asynchronous. This keeps opening a
    new Angebot/Rechnung independent of the size of the imported price library.
    """
    rows = list(
        m.CatalogItem.objects.filter(organization=org, active=True)
        .filter(Q(sales_price__gt=0) | Q(purchase_price__gt=0))
        .only("id", "code", "name", "description", "unit", "kind", "purchase_price", "sales_price")
        .order_by("name")[: max(1, min(int(limit or 18), 40))]
    )
    for item in rows:
        local_price = _money(item.sales_price) if _money(item.sales_price) > 0 else _money(item.purchase_price)
        item.effective_sales_price = local_price
        item.effective_price_source = "A+Bau-Vorlage"
        item.effective_price_source_kind = "A+Bau"
        item.effective_price_reference_code = item.code or ""
        item.effective_price_match_kind = "local"
    return rows


@login_required
@require_http_methods(["GET", "POST"])
def quote_editor(request, pk=None):
    org = _org(request)
    quote = get_object_or_404(m.Quote, pk=pk, organization=org) if pk else None
    initial = {}
    if request.GET.get("project"):
        initial["project"] = request.GET.get("project")
    form = QuoteForm(request.POST or None, instance=quote, organization=org, initial=initial)
    if request.method == "POST" and form.is_valid():
        with transaction.atomic():
            obj = form.save(commit=False)
            obj.organization = org
            obj.created_by = obj.created_by or request.user
            if not obj.number:
                obj.number = _unique_number(m.Quote, org, "A")
            if request.POST.get("action") == "send":
                obj.status = "sent"
                obj.sent_at = timezone.now()
            obj.outro_text = (request.POST.get("closing_text") or "").strip()
            obj.save()
            settings = _save_document_settings(obj, request, "quote")
            obj.discount_percent = settings.discount_value if settings.discount_type == "percent" else Decimal("0")
            obj.save(update_fields=["discount_percent", "outro_text", "updated_at"])
            _save_commercial_items(obj, request, "quote", settings)
        messages.success(request, "Angebot gespeichert.")
        return redirect("next-quote-edit", pk=obj.pk)
    settings = _commercial_settings(quote)
    catalog = _fast_catalog_preview(org)
    return render(request, "rebuild/document_editor.html", {
        "form": form, "document": quote, "items": _items_for_editor(quote), "catalog": catalog,
        "kind": "quote", "totals": _quote_total(quote) if quote else None, "commercial": settings,
    })


@login_required
def invoice_list(request):
    org = _org(request)
    invoices = m.Invoice.objects.filter(organization=org).select_related("project", "project__customer").prefetch_related("items", "payments").order_by("-created_at")[:250]
    rows = [{"invoice": invoice, "total": _invoice_total(invoice)} for invoice in invoices]
    return render(request, "rebuild/invoices.html", {"rows": rows})


def _save_invoice_items(invoice, request):
    settings = _commercial_settings(invoice)
    if settings is None:
        settings = m.CommercialDocumentSettings.objects.create(organization=invoice.organization, invoice=invoice)
    _save_commercial_items(invoice, request, "invoice", settings)


@login_required
@require_http_methods(["GET", "POST"])
def invoice_editor(request, pk=None):
    org = _org(request)
    invoice = get_object_or_404(m.Invoice, pk=pk, organization=org) if pk else None
    if invoice is not None and is_finalized(invoice) and request.method == "POST":
        messages.error(request, "Finalisierte Rechnungen sind unveränderbar. Bitte Storno oder Korrektur verwenden.")
        return redirect("next-invoice-edit", pk=invoice.pk)
    initial = {}
    if request.GET.get("project"):
        initial["project"] = request.GET.get("project")
    if request.GET.get("quote"):
        initial["quote"] = request.GET.get("quote")
    form = InvoiceForm(request.POST or None, instance=invoice, organization=org, initial=initial)
    if request.method == "POST" and form.is_valid():
        action = request.POST.get("action") or "save"
        try:
            with transaction.atomic():
                obj = form.save(commit=False)
                obj.organization = org
                obj.created_by = obj.created_by or request.user
                # A draft deliberately has no fiscal invoice number. The number is
                # allocated transactionally only during finalization.
                if not obj.pk:
                    obj.number = ""
                    obj.status = "draft"
                due_days = max(0, min(365, int(_money(request.POST.get("payment_due_days") or 14))))
                obj.due_date = obj.issue_date + timedelta(days=due_days)
                obj.outro_text = (request.POST.get("closing_text") or "").strip()
                obj.save()
                settings = _save_document_settings(obj, request, "invoice")
                _save_commercial_items(obj, request, "invoice", settings)
                save_customer_profile(obj, request.POST)
                if action == "finalize":
                    finalize_invoice(obj, user=request.user, request=request)
        except ComplianceError as exc:
            messages.error(request, "Rechnung kann nicht finalisiert werden: " + str(exc))
            return render(request, "rebuild/document_editor.html", {
                "form": form, "document": invoice, "items": _items_for_editor(invoice),
                "catalog": _fast_catalog_preview(org),
                "kind": "invoice", "totals": _invoice_total(invoice) if invoice else None, "commercial": _commercial_settings(invoice),
                "invoice_compliance": get_compliance(invoice),
            })
        if action == "finalize":
            messages.success(request, "Rechnung finalisiert, nummeriert, gehasht und archiviert.")
        else:
            messages.success(request, "Rechnungsentwurf gespeichert.")
        return redirect("next-invoice-edit", pk=obj.pk)
    settings = _commercial_settings(invoice)
    return render(request, "rebuild/document_editor.html", {
        "form": form, "document": invoice, "items": _items_for_editor(invoice),
        "catalog": _fast_catalog_preview(org),
        "kind": "invoice", "totals": _invoice_total(invoice) if invoice else None, "commercial": settings,
        "invoice_compliance": get_compliance(invoice),
    })


@login_required
@require_POST
def invoice_payment(request, pk):
    org = _org(request)
    invoice = get_object_or_404(m.Invoice, pk=pk, organization=org)
    amount = _money(request.POST.get("amount"))
    if amount <= 0:
        messages.error(request, "Bitte einen gültigen Betrag eingeben.")
        return redirect("next-invoice-edit", pk=invoice.pk)
    m.Payment.objects.create(invoice=invoice, amount=amount, paid_at=timezone.localdate(), method=request.POST.get("method") or "Überweisung", reference=request.POST.get("reference") or "", recorded_by=request.user)
    totals = _invoice_total(invoice)
    invoice.status = "paid" if totals["open"] <= 0 else "partial"
    invoice.save(update_fields=["status", "updated_at"])
    messages.success(request, "Zahlung erfasst.")
    return redirect("next-invoice-edit", pk=invoice.pk)


@login_required
def finance_dashboard(request):
    if _is_field_user(request):
        messages.error(request, "Die Finanzübersicht ist nur für Büro, Projektleitung und Buchhaltung verfügbar.")
        return redirect("next-dashboard")
    org = _org(request)
    invoices = m.Invoice.objects.filter(organization=org).exclude(status="cancelled").select_related("project", "project__customer").prefetch_related("items", "items__catalog_item", "items__commercial_meta", "payments")
    quotes = m.Quote.objects.filter(organization=org).exclude(status="rejected").prefetch_related("items", "items__catalog_item", "items__commercial_meta")
    date_from = (request.GET.get("date_from") or "").strip()
    date_to = (request.GET.get("date_to") or "").strip()
    project_filter = (request.GET.get("project") or "").strip()
    try:
        if date_from:
            invoices = invoices.filter(issue_date__gte=timezone.datetime.strptime(date_from, "%Y-%m-%d").date())
            quotes = quotes.filter(issue_date__gte=timezone.datetime.strptime(date_from, "%Y-%m-%d").date())
    except ValueError:
        date_from = ""
    try:
        if date_to:
            invoices = invoices.filter(issue_date__lte=timezone.datetime.strptime(date_to, "%Y-%m-%d").date())
            quotes = quotes.filter(issue_date__lte=timezone.datetime.strptime(date_to, "%Y-%m-%d").date())
    except ValueError:
        date_to = ""
    if project_filter.isdigit():
        invoices = invoices.filter(project_id=int(project_filter))
        quotes = quotes.filter(project_id=int(project_filter))
    invoice_rows = []
    revenue = cost = gross = paid = open_amount = Decimal("0")
    by_project = {}
    for invoice in invoices.order_by("-issue_date", "-created_at"):
        totals = _invoice_total(invoice)
        revenue += totals["net"]; cost += totals["cost"]; gross += totals["gross"]; paid += totals["paid"]; open_amount += totals["open"]
        invoice_rows.append({"invoice": invoice, "total": totals})
        bucket = by_project.setdefault(invoice.project_id, {"project": invoice.project, "revenue": Decimal("0"), "cost": Decimal("0"), "margin": Decimal("0"), "gross": Decimal("0"), "open": Decimal("0")})
        bucket["revenue"] += totals["net"]; bucket["cost"] += totals["cost"]; bucket["gross"] += totals["gross"]; bucket["open"] += totals["open"]
    project_rows = []
    for bucket in by_project.values():
        bucket["margin"] = bucket["revenue"] - bucket["cost"]
        bucket["margin_percent"] = (bucket["margin"] / bucket["revenue"] * Decimal("100")) if bucket["revenue"] else Decimal("0")
        project_rows.append(bucket)
    project_rows.sort(key=lambda row: row["revenue"], reverse=True)
    margin = revenue - cost
    quote_volume = sum((_quote_total(q)["net"] for q in quotes), Decimal("0"))
    projects = m.Project.objects.filter(organization=org, archived=False).order_by("-updated_at")[:300]
    return render(request, "rebuild/finance_dashboard.html", {
        "revenue": revenue, "cost": cost, "margin": margin,
        "margin_percent": (margin / revenue * Decimal("100")) if revenue else Decimal("0"),
        "gross": gross, "paid": paid, "open_amount": open_amount, "quote_volume": quote_volume,
        "invoice_rows": invoice_rows[:100], "project_rows": project_rows[:100], "projects": projects,
        "date_from": date_from, "date_to": date_to, "project_filter": project_filter,
    })


@login_required
def time_overview(request):
    org = _org(request); employee = _employee(request, org); field_user = _is_field_user(request)
    scoped = m.TimeEntry.objects.filter(organization=org).select_related("employee", "project")
    if field_user and employee: scoped = scoped.filter(employee=employee)
    employee_filter=(request.GET.get("employee") or "").strip(); project_filter=(request.GET.get("project") or "").strip(); date_from=(request.GET.get("date_from") or "").strip(); date_to=(request.GET.get("date_to") or "").strip()
    if not field_user and employee_filter.isdigit(): scoped=scoped.filter(employee_id=int(employee_filter))
    if project_filter.isdigit(): scoped=scoped.filter(project_id=int(project_filter))
    today=timezone.localdate(); week_start=today-timedelta(days=today.weekday()); now=timezone.now()
    def mins(e):
        raw=max(0,int(((e.ended_at or now)-e.started_at).total_seconds()//60)); return max(0,raw-int(e.break_minutes or 0))
    today_minutes=sum(mins(e) for e in scoped.filter(started_at__date=today)); week_minutes=sum(mins(e) for e in scoped.filter(started_at__date__gte=week_start,started_at__date__lte=today)); running_count=scoped.filter(ended_at__isnull=True).count()
    entries=scoped
    try:
        if date_from: entries=entries.filter(started_at__date__gte=timezone.datetime.strptime(date_from,"%Y-%m-%d").date())
    except ValueError: date_from=""
    try:
        if date_to: entries=entries.filter(started_at__date__lte=timezone.datetime.strptime(date_to,"%Y-%m-%d").date())
    except ValueError: date_to=""
    entries=list(entries.order_by("-started_at")[:150])
    for e in entries:
        mnt=mins(e); e.ui_duration=f"{mnt//60:d}:{mnt%60:02d} h"
    employees=m.Employee.objects.filter(organization=org,active=True).order_by("last_name","first_name") if not field_user else m.Employee.objects.none(); projects=m.Project.objects.filter(organization=org,archived=False).order_by("-updated_at")[:250]
    return render(request,"rebuild/time_overview.html",{"entries":entries,"employee":employee,"field_user":field_user,"employees":employees,"projects":projects,"employee_filter":employee_filter,"project_filter":project_filter,"date_from":date_from,"date_to":date_to,"today_hours":f"{today_minutes//60:d}:{today_minutes%60:02d}","week_hours":f"{week_minutes//60:d}:{week_minutes%60:02d}","running_count":running_count,"week_start":week_start})


@login_required
@require_http_methods(["GET","POST"])
def time_entry_edit(request, pk):
    org=_org(request)
    if _is_field_user(request): messages.error(request,"Arbeitszeitkorrekturen sind nur für Büro/Leitung verfügbar."); return redirect("next-time")
    entry=get_object_or_404(m.TimeEntry.objects.select_related("employee","project"),pk=pk,organization=org); error=""
    if request.method=="POST":
        try:
            start=timezone.make_aware(timezone.datetime.strptime((request.POST.get("started_at") or "").strip(),"%Y-%m-%dT%H:%M"),timezone.get_current_timezone()); end_raw=(request.POST.get("ended_at") or "").strip(); end=timezone.make_aware(timezone.datetime.strptime(end_raw,"%Y-%m-%dT%H:%M"),timezone.get_current_timezone()) if end_raw else None; pause=int((request.POST.get("break_minutes") or "0").strip() or 0)
            if end and end<start: raise ValueError("Ende liegt vor dem Start.")
            if pause<0 or pause>1440: raise ValueError("Pause ist ungültig.")
            entry.started_at=start; entry.ended_at=end; entry.break_minutes=pause; entry.description=(request.POST.get("description") or "")[:500]; entry.approved=request.POST.get("approved")=="1"; entry.save(update_fields=["started_at","ended_at","break_minutes","description","approved","updated_at"]); messages.success(request,"Arbeitszeit wurde korrigiert."); return redirect("next-time")
        except (ValueError,TypeError) as exc: error=str(exc) or "Bitte die Zeitangaben prüfen."
    return render(request,"rebuild/time_entry_form.html",{"entry":entry,"error":error})


@login_required
def migration_import(request):
    org = _org(request)
    summary = None
    errors = []
    if request.method == "POST" and request.FILES.get("file"):
        file = request.FILES["file"]
        kind = request.POST.get("kind") or "customers"
        try:
            rows = _read_table(file)
            summary = _import_tooltime_rows(org, request.user, kind, rows)
            messages.success(request, f"ToolTime-Import abgeschlossen: {summary['created']} neu, {summary['updated']} aktualisiert.")
        except Exception as exc:
            errors.append(str(exc))
    return render(request, "rebuild/migration.html", {"summary": summary, "errors": errors})


def _read_table(upload):
    name = upload.name.lower()
    raw = upload.read()
    if name.endswith(".xlsx"):
        try:
            import openpyxl
        except ImportError as exc:
            raise ValueError("XLSX-Unterstützung fehlt. Bitte CSV exportieren oder openpyxl installieren.") from exc
        book = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        sheet = book.active
        values = list(sheet.iter_rows(values_only=True))
        if not values:
            return []
        headers = [str(value or "").strip() for value in values[0]]
        return [dict(zip(headers, row)) for row in values[1:] if any(value not in (None, "") for value in row)]
    text = raw.decode("utf-8-sig", errors="replace")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t,")
    except csv.Error:
        dialect = csv.excel
        dialect.delimiter = ";"
    return list(csv.DictReader(io.StringIO(text), dialect=dialect))


def _norm(row):
    return {re.sub(r"[^a-z0-9]", "", str(key).lower().replace("ä", "ae").replace("ö", "oe").replace("ü", "ue").replace("ß", "ss")): ("" if value is None else str(value).strip()) for key, value in row.items()}


def _pick(row, *aliases):
    for alias in aliases:
        key = re.sub(r"[^a-z0-9]", "", alias.lower().replace("ä", "ae").replace("ö", "oe").replace("ü", "ue").replace("ß", "ss"))
        if row.get(key):
            return row[key]
    return ""


def _import_tooltime_rows(org, user, kind, rows):
    created = updated = skipped = 0
    with transaction.atomic():
        if kind == "customers":
            for source in rows:
                row = _norm(source)
                number = _pick(row, "Kundennummer", "Customer Number", "Nummer")
                email = _pick(row, "E-Mail", "Email")
                company = _pick(row, "Firma", "Unternehmen", "Company")
                first = _pick(row, "Vorname", "First Name")
                last = _pick(row, "Nachname", "Last Name")
                query = m.Customer.objects.filter(organization=org)
                customer = query.filter(number=number).first() if number else None
                if customer is None and email:
                    customer = query.filter(email__iexact=email).first()
                values = {
                    "company": company, "first_name": first, "last_name": last, "email": email,
                    "phone": _pick(row, "Telefon", "Phone"), "mobile": _pick(row, "Mobil", "Mobile"),
                    "street": _pick(row, "Straße", "Strasse", "Street"), "postal_code": _pick(row, "PLZ", "Postleitzahl", "ZIP"),
                    "city": _pick(row, "Ort", "Stadt", "City"), "country": _pick(row, "Land", "Country") or "DE",
                }
                if customer:
                    for key, value in values.items():
                        if value:
                            setattr(customer, key, value)
                    customer.save()
                    updated += 1
                else:
                    m.Customer.objects.create(organization=org, number=number or _unique_number(m.Customer, org, "K"), **values)
                    created += 1
        elif kind in {"quotes", "invoices"}:
            model = m.Quote if kind == "quotes" else m.Invoice
            prefix = "A" if kind == "quotes" else "R"
            grouped = {}
            for source in rows:
                row = _norm(source)
                number = _pick(row, "Angebotsnummer" if kind == "quotes" else "Rechnungsnummer", "Nummer", "Number")
                if not number:
                    skipped += 1
                    continue
                grouped.setdefault(number, []).append(row)
            for number, group in grouped.items():
                first = group[0]
                customer_number = _pick(first, "Kundennummer", "Customer Number")
                customer_name = _pick(first, "Kunde", "Kundenname", "Customer")
                customer = m.Customer.objects.filter(organization=org, number=customer_number).first() if customer_number else None
                if customer is None and customer_name:
                    customer = m.Customer.objects.filter(organization=org).filter(Q(company__iexact=customer_name) | Q(last_name__iexact=customer_name)).first()
                if customer is None:
                    customer = m.Customer.objects.create(organization=org, number=_unique_number(m.Customer, org, "K"), company=customer_name or "ToolTime Import")
                project_title = _pick(first, "Projekt", "Projekttitel", "Project") or f"ToolTime Import {number}"
                project = m.Project.objects.filter(organization=org, customer=customer, title=project_title).first()
                if project is None:
                    project = m.Project.objects.create(organization=org, customer=customer, number=_unique_number(m.Project, org, "P"), title=project_title, status="invoiced" if kind == "invoices" else "quoted")
                obj = model.objects.filter(organization=org, number=number).first()
                if obj:
                    updated += 1
                else:
                    kwargs = {"organization": org, "number": number or _unique_number(model, org, prefix), "project": project, "created_by": user}
                    if kind == "invoices":
                        kwargs["due_date"] = timezone.localdate() + timedelta(days=14)
                    obj = model.objects.create(**kwargs)
                    created += 1
                item_model = m.QuoteItem if kind == "quotes" else m.InvoiceItem
                parent_field = "quote" if kind == "quotes" else "invoice"
                if not getattr(obj, "items").exists():
                    for position, row in enumerate(group, 1):
                        description = _pick(row, "Positionsbezeichnung", "Beschreibung", "Description", "Leistung", "Artikel")
                        if not description:
                            continue
                        item_model.objects.create(**{
                            parent_field: obj, "position": position, "description": description,
                            "quantity": _money(_pick(row, "Menge", "Quantity") or 1),
                            "unit": _pick(row, "Einheit", "Unit") or "Stk.",
                            "unit_price": _money(_pick(row, "Einzelpreis", "Preis", "Unit Price", "Netto")),
                            "tax_rate": _money(_pick(row, "MwSt", "Steuer", "Tax") or 19),
                        })
        elif kind == "time":
            employee = m.Employee.objects.filter(organization=org, user=user).first() or m.Employee.objects.filter(organization=org).first()
            if employee is None:
                raise ValueError("Für den Zeitimport muss mindestens ein Mitarbeiter existieren.")
            for source in rows:
                row = _norm(source)
                project_number = _pick(row, "Projektnummer", "Project Number")
                project = m.Project.objects.filter(organization=org, number=project_number).first() if project_number else None
                if project is None:
                    skipped += 1
                    continue
                date_text = _pick(row, "Datum", "Date")
                start_text = _pick(row, "Start", "Von", "Startzeit")
                end_text = _pick(row, "Ende", "Bis", "Endzeit")
                try:
                    date = timezone.datetime.strptime(date_text, "%d.%m.%Y").date() if "." in date_text else timezone.datetime.fromisoformat(date_text).date()
                    start_time = timezone.datetime.strptime(start_text, "%H:%M").time()
                    end_time = timezone.datetime.strptime(end_text, "%H:%M").time()
                except Exception:
                    skipped += 1
                    continue
                started = timezone.make_aware(timezone.datetime.combine(date, start_time))
                ended = timezone.make_aware(timezone.datetime.combine(date, end_time))
                m.TimeEntry.objects.create(organization=org, employee=employee, project=project, started_at=started, ended_at=ended, break_minutes=int(_money(_pick(row, "Pause", "Pausenminuten") or 0)), description=_pick(row, "Beschreibung", "Description"))
                created += 1
        else:
            raise ValueError("Unbekannter Importtyp.")
    return {"created": created, "updated": updated, "skipped": skipped, "rows": len(rows)}


@login_required
def settings_page(request):
    org = _org(request)
    integrations = m.IntegrationConfig.objects.filter(organization=org).order_by("provider")
    return render(request, "rebuild/settings.html", {"organization": org, "integrations": integrations})

# A+BAU INVOICE COMPLIANCE CATALOG PERFORMANCE FIX 2026-08-20
